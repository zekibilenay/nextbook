#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Next Book — personal reading-priority list (Windows desktop app)

* Searches Open Library for title, author, genre and page count only.
* Each book becomes one row. Scored rows get more intense the higher the
  score; unscored rows use normal text.
* Scoring criteria and their weights are fully user-defined.
* 6 languages, 6 colour themes, font choice, Excel import/export.

Requires : Python 3.9+ (tkinter ships with the Windows installer)
Optional : pip install openpyxl   (Excel import/export)
Run      : python next_book.py
"""

import base64
import ctypes
import json
import math
import os
import queue
import re
import sys
import threading
import urllib.error
import urllib.parse
import urllib.request
import struct
import uuid
import warnings
import zlib
import webbrowser
import datetime
import tkinter as tk
from pathlib import Path
from tkinter import ttk, messagebox, filedialog, simpledialog
import tkinter.font as tkfont

APP_NAME = "Next Book"
APP_VERSION = "1.8.1"
APP_ID = "nextbook"          # duyuru sunucusunda bu uygulamayı tanımlayan kimlik
REPO_URL = "https://github.com/zekibilenay/nextbook"
def _default_data_dir():
    """Windows: %APPDATA%\\NextBook · macOS: ~/Library/Application Support/NextBook · Linux: ~/.local/share/NextBook"""
    if sys.platform == "win32":
        base = os.environ.get("APPDATA") or Path.home()
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    else:
        base = os.environ.get("XDG_DATA_HOME") or Path.home() / ".local" / "share"
    return Path(base) / "NextBook"


DATA_DIR = _default_data_dir()
# macOS'ta sağ tık <Button-2>, orta tık <Button-3>; Windows / Linux'ta tersi
RCLICK, MCLICK = ("<Button-2>", "<Button-3>") if sys.platform == "darwin" else ("<Button-3>", "<Button-2>")
RCLICK_PRESS = RCLICK.replace("Button", "ButtonPress")
DATA_FILE = DATA_DIR / "next_book.json"
SETTINGS_FILE = DATA_DIR / "settings.json"     # dil / tema / yazı tipi (veriden ayrı)

# Duyuru kaynağı: tüm uygulamaların ortak kullandığı tek bir JSON dosyası (bkz. announcements.example.json).
# Sadece HTTPS kabul edilir. Boş bırakırsan duyuru özelliği tamamen kapanır.
ANNOUNCE_URL = os.environ.get("NEXTBOOK_ANNOUNCE_URL",
                              "https://zekibilenay.github.io/announcements/feed.json")
ANNOUNCE_REFRESH_MS = 6 * 60 * 60 * 1000       # açık kaldığı sürece 6 saatte bir yeniden sor
ANNOUNCE_ROTATE_MS = 10_000                     # birden çok duyuru varsa değişme süresi

# Yazı tipi seçicide yalnızca bilgisayarlarda yaygın, okunaklı bu yazı tipleri sunulur
# (bilgisayarda yüklü olmayanlar listeden otomatik çıkar).
FONT_CHOICES = ("Segoe UI", "Calibri", "Arial", "Verdana", "Tahoma", "Trebuchet MS",
                "Georgia", "Cambria", "Times New Roman", "Consolas")

DEFAULT_WEIGHTS = [0.30, 0.30, 0.15, 0.10, 0.15]
DEFAULT_SCALE_MAX = 10
STATUSES = ("todo", "reading", "paused", "done")   # okunacak · okuyorum · yarım bıraktım · okudum
NEW_IID = "__new__"                              # tablonun en altındaki hep boş satır
MAX_LISTS = 20                                  # en fazla kaç sekme (liste) açılabilir


# ============================================================================
# ÇOK DİLLİLİK
# ============================================================================

LANGS = [("tr", "Türkçe"), ("en", "English"), ("ru", "Русский"),
         ("de", "Deutsch"), ("fr", "Français"), ("zh", "中文")]

I18N = {
"tr": {
 "lbl_title": "Kitap adı", "lbl_author": "Yazar",
 "btn_search": "Ara", "btn_clear": "Temizle", "btn_manual": "+ Elle ekle",
 "results_title": " Arama sonuçları — eklemek için çift tıkla ",
 "btn_add_to_list": "Listeye ekle", "list_title": " Okuma listem ",
 "btn_settings": "⚙ Puanlama ayarları", "btn_sort_score": "Puana göre sırala",
 "btn_delete": "Seçileni sil", "btn_import": "Excel'den içe aktar", "btn_export": "Excel'e aktar",
 "hint": "Hücreye çift tıkla → düzenle · Enter = aşağı, Tab = sağa, Esc = vazgeç, boş bırak = puanı sil · "
         "Renk ne kadar yoğunsa skor o kadar yüksek · Başlığa tıklayarak sırala · Satıra sağ tıkla: durum, sil",
 "col_title": "Kitap Adı", "col_author": "Yazar", "col_genre": "Tür", "col_pages": "Sayfa",
 "col_year": "Yıl", "col_score": "Öncelik Skoru",
 "status_counts": "{total} kitap · {scored} puanlandı · {pending} puan bekliyor",
 "searching": "Aranıyor…", "enter_query": "Bir kitap adı veya yazar yazın.",
 "no_results": "Sonuç bulunamadı. Yazımı değiştirmeyi ya da “+ Elle ekle”yi deneyin.",
 "results_found": "{n} sonuç bulundu. Eklemek için çift tıklayın.",
 "pick_result": "Önce bir sonuç seçin.",
 "err_network": "Open Library'ye ulaşılamadı. İnternet bağlantınızı kontrol edin.",
 "err_search": "Arama hatası: {err}",
 "dup_book": "“{title}” zaten listenizde.", "ask_title": "Kitap adı:",
 "confirm_delete": "Seçili {n} kitap listeden silinsin mi?",
 "bad_pages": "Sayfa sayısı bir sayı olmalı.",
 "bad_score": "Geçersiz puan: 1 ile {max} arasında bir sayı girin.",
 "need_openpyxl": "Excel desteği için openpyxl gerekli.\n\nKomut satırında şunu çalıştırın:\n\npip install openpyxl",
 "pick_excel": "Excel dosyası seç",
 "confirm_import": "Mevcut listeniz silinip Excel dosyasındaki kitaplar ve ölçütler yüklenecek.\n"
                   "(Önce “Excel'e aktar” ile yedek almak isteyebilirsiniz.)\n\nDevam edilsin mi?",
 "import_failed": "Excel okunamadı:\n{err}",
 "import_done": "{books} kitap ve {crit} ölçüt içe aktarıldı.",
 "save_excel": "Excel olarak kaydet",
 "export_locked": "Dosya yazılamadı. Excel'de açıksa kapatıp tekrar deneyin.",
 "export_failed": "Dışa aktarılamadı:\n{err}", "saved_to": "Kaydedildi: {path}",
 "save_failed": "Kaydedilemedi:\n{err}",
 "xl_sheet_main": "Okuma Önceliklendirme", "xl_sheet_settings": "Ayarlar",
 "xl_weights_head": "Ağırlıklar", "xl_weight": "Ağırlık", "xl_total": "Toplam",
 "xl_note": "Ağırlıklar birbirine oranlanır; toplamın 1 olması şart değil.",
 "st_title": "Puanlama ayarları", "st_heading": "Ölçütlerini ve ağırlıklarını kendin belirle.",
 "st_desc": "Ağırlığı yüksek ölçüt, nihai skoru daha çok etkiler. Ağırlıklar birbirine oranlanır; "
            "toplamın 1 olması gerekmez. Boş bırakılan puan hesaba katılmaz.",
 "st_name": "Ölçüt adı", "st_weight": "Ağırlık", "st_add": "+ Ölçüt ekle",
 "st_scale": "Puan üst sınırı:", "st_total": "Ağırlık toplamı: {total}",
 "btn_cancel": "İptal", "btn_save": "Kaydet",
 "st_need_one": "En az bir ölçüt olmalı.", "st_name_required": "Tüm ölçütlerin bir adı olmalı.",
 "st_bad_weight": "“{name}” için ağırlık 0 veya daha büyük bir sayı olmalı.",
 "st_dup_name": "İki ölçüt aynı ada sahip olamaz.",
 "st_bad_scale": "Puan üst sınırı 2 ile 10 arasında olmalı.",
 "font_title": "Yazı tipi", "font_family": "Yazı tipi", "font_size": "Boyut",
 "font_preview": "Okuma sırası belli olsun — Aa Bb 123", "font_reset": "Varsayılan", "btn_apply": "Uygula",
 "about_title": "Hakkında",
 "crit_1": "İlgi Düzeyi", "crit_2": "Kişisel Katkı", "crit_3": "Okuma Kolaylığı",
 "crit_4": "Sosyal Bağlam", "crit_5": "Uzun Vadeli / Referans Değeri",
 "ann_new_version": "Yeni sürüm yayınlandı: {latest} (sizdeki: {current})", "ann_open_q": "Bağlantı tarayıcıda açılsın mı?\n{url}",
 "ann_board_title": "Duyurular",
 "ann_empty": "Şu an duyuru yok.",
 "ann_type_update": "Güncelleme",
 "ann_type_new_app": "Yeni uygulama",
 "ann_type_recommended": "Önerilen uygulama",
 "ann_type_info": "Bilgi",
 "ann_open": "Bağlantıyı aç",
 "btn_close": "Kapat",
 "about_hello": "Merhaba! Ben Zeki. Next Book'u, okuyacağın kitapları kendi ölçütlerine göre puanlayıp sana doğru okuma sırasını göstermesi için geliştiriyorum.",
 "about_privacy": "Verilerin yalnızca bilgisayarında kalır; hesap gerektirmez. İnternet yalnızca kitap araması (Open Library) ve duyurular için kullanılır.",
 "about_data": "Veri klasörü",
        "data_recovered": "Veri dosyan okunamadı (bozulmuş olabilir). Yeni bir liste ile başlandı; eski dosya şuraya yedeklendi:\n{path}",
},
"en": {
 "lbl_title": "Book title", "lbl_author": "Author",
 "btn_search": "Search", "btn_clear": "Clear", "btn_manual": "+ Add manually",
 "results_title": " Search results — double-click to add ",
 "btn_add_to_list": "Add to list", "list_title": " My reading list ",
 "btn_settings": "⚙ Scoring settings", "btn_sort_score": "Sort by score",
 "btn_delete": "Delete selected", "btn_import": "Import from Excel", "btn_export": "Export to Excel",
 "hint": "Double-click a cell to edit · Enter = down, Tab = right, Esc = cancel, empty = clear score · "
         "The more intense the colour, the higher the score · Click a header to sort · Right-click a book: status, remove",
 "col_title": "Title", "col_author": "Author", "col_genre": "Genre", "col_pages": "Pages",
 "col_year": "Year", "col_score": "Priority Score",
 "status_counts": "{total} books · {scored} scored · {pending} waiting for a score",
 "searching": "Searching…", "enter_query": "Type a book title or an author.",
 "no_results": "No results. Try a different spelling or “+ Add manually”.",
 "results_found": "{n} results found. Double-click to add.",
 "pick_result": "Select a result first.",
 "err_network": "Could not reach Open Library. Please check your internet connection.",
 "err_search": "Search error: {err}",
 "dup_book": "“{title}” is already in your list.", "ask_title": "Book title:",
 "confirm_delete": "Delete the {n} selected book(s) from your list?",
 "bad_pages": "Page count must be a number.",
 "bad_score": "Invalid score: enter a number between 1 and {max}.",
 "need_openpyxl": "Excel support needs openpyxl.\n\nRun this in a command prompt:\n\npip install openpyxl",
 "pick_excel": "Choose an Excel file",
 "confirm_import": "Your current list will be replaced by the books and criteria from the Excel file.\n"
                   "(You may want to back up first with “Export to Excel”.)\n\nContinue?",
 "import_failed": "Could not read the Excel file:\n{err}",
 "import_done": "Imported {books} books and {crit} criteria.",
 "save_excel": "Save as Excel",
 "export_locked": "Could not write the file. If it is open in Excel, close it and try again.",
 "export_failed": "Could not export:\n{err}", "saved_to": "Saved: {path}",
 "save_failed": "Could not save:\n{err}",
 "xl_sheet_main": "Reading Priorities", "xl_sheet_settings": "Settings",
 "xl_weights_head": "Weights", "xl_weight": "Weight", "xl_total": "Total",
 "xl_note": "Weights are relative to each other; they do not need to add up to 1.",
 "st_title": "Scoring settings", "st_heading": "Define your own criteria and weights.",
 "st_desc": "A criterion with a higher weight affects the final score more. Weights are relative to each "
            "other and do not need to add up to 1. Empty scores are ignored.",
 "st_name": "Criterion", "st_weight": "Weight", "st_add": "+ Add criterion",
 "st_scale": "Maximum score:", "st_total": "Total weight: {total}",
 "btn_cancel": "Cancel", "btn_save": "Save",
 "st_need_one": "There must be at least one criterion.", "st_name_required": "Every criterion needs a name.",
 "st_bad_weight": "The weight for “{name}” must be 0 or greater.",
 "st_dup_name": "Two criteria cannot have the same name.",
 "st_bad_scale": "The maximum score must be between 2 and 10.",
 "font_title": "Font", "font_family": "Font", "font_size": "Size",
 "font_preview": "Know what to read next — Aa Bb 123", "font_reset": "Default", "btn_apply": "Apply",
 "about_title": "About",
 "crit_1": "Interest", "crit_2": "Personal Value", "crit_3": "Ease of Reading",
 "crit_4": "Social Context", "crit_5": "Long-term / Reference Value",
 "ann_new_version": "A new version is out: {latest} (you have {current})", "ann_open_q": "Open this link in your browser?\n{url}",
 "ann_board_title": "Announcements",
 "ann_empty": "No announcements right now.",
 "ann_type_update": "Update",
 "ann_type_new_app": "New app",
 "ann_type_recommended": "Recommended app",
 "ann_type_info": "Info",
 "ann_open": "Open link",
 "btn_close": "Close",
 "about_hello": "Hi! I'm Zeki. I'm building Next Book to help you score the books you want to read by your own criteria and see what to read next.",
 "about_privacy": "Your data stays on your computer and no account is needed. The internet is used only for book search (Open Library) and announcements.",
 "about_data": "Data folder",
        "data_recovered": "Your data file could not be read (it may be corrupted). A new list was started; the old file was backed up to:\n{path}",
},
"ru": {
 "lbl_title": "Название книги", "lbl_author": "Автор",
 "btn_search": "Найти", "btn_clear": "Очистить", "btn_manual": "+ Добавить вручную",
 "results_title": " Результаты поиска — двойной щелчок, чтобы добавить ",
 "btn_add_to_list": "Добавить в список", "list_title": " Мой список чтения ",
 "btn_settings": "⚙ Настройки оценки", "btn_sort_score": "Сортировать по оценке",
 "btn_delete": "Удалить выбранное", "btn_import": "Импорт из Excel", "btn_export": "Экспорт в Excel",
 "hint": "Двойной щелчок по ячейке — редактировать · Enter — вниз, Tab — вправо, Esc — отмена, пусто — стереть оценку · "
         "Чем насыщеннее цвет, тем выше оценка · Щёлкните заголовок для сортировки · ПКМ по книге: статус, удалить",
 "col_title": "Название", "col_author": "Автор", "col_genre": "Жанр", "col_pages": "Стр.",
 "col_year": "Год", "col_score": "Приоритет",
 "status_counts": "Книг: {total} · оценено: {scored} · ждут оценки: {pending}",
 "searching": "Поиск…", "enter_query": "Введите название книги или автора.",
 "no_results": "Ничего не найдено. Измените запрос или нажмите «+ Добавить вручную».",
 "results_found": "Найдено результатов: {n}. Дважды щёлкните, чтобы добавить.",
 "pick_result": "Сначала выберите результат.",
 "err_network": "Не удалось подключиться к Open Library. Проверьте подключение к интернету.",
 "err_search": "Ошибка поиска: {err}",
 "dup_book": "«{title}» уже есть в вашем списке.", "ask_title": "Название книги:",
 "confirm_delete": "Удалить выбранные книги ({n}) из списка?",
 "bad_pages": "Количество страниц должно быть числом.",
 "bad_score": "Недопустимая оценка: введите число от 1 до {max}.",
 "need_openpyxl": "Для работы с Excel нужен openpyxl.\n\nВыполните в командной строке:\n\npip install openpyxl",
 "pick_excel": "Выберите файл Excel",
 "confirm_import": "Текущий список будет заменён книгами и критериями из файла Excel.\n"
                   "(Сначала можно сделать резервную копию через «Экспорт в Excel».)\n\nПродолжить?",
 "import_failed": "Не удалось прочитать файл Excel:\n{err}",
 "import_done": "Импортировано книг: {books}, критериев: {crit}.",
 "save_excel": "Сохранить как Excel",
 "export_locked": "Не удалось записать файл. Если он открыт в Excel, закройте его и повторите.",
 "export_failed": "Не удалось экспортировать:\n{err}", "saved_to": "Сохранено: {path}",
 "save_failed": "Не удалось сохранить:\n{err}",
 "xl_sheet_main": "Приоритеты чтения", "xl_sheet_settings": "Настройки",
 "xl_weights_head": "Веса", "xl_weight": "Вес", "xl_total": "Итого",
 "xl_note": "Веса соотносятся друг с другом; их сумма не обязана равняться 1.",
 "st_title": "Настройки оценки", "st_heading": "Задайте свои критерии и их веса.",
 "st_desc": "Критерий с большим весом сильнее влияет на итоговую оценку. Веса соотносятся друг с другом, "
            "их сумма не обязана равняться 1. Пустые оценки не учитываются.",
 "st_name": "Критерий", "st_weight": "Вес", "st_add": "+ Добавить критерий",
 "st_scale": "Максимальная оценка:", "st_total": "Сумма весов: {total}",
 "btn_cancel": "Отмена", "btn_save": "Сохранить",
 "st_need_one": "Должен остаться хотя бы один критерий.", "st_name_required": "У каждого критерия должно быть название.",
 "st_bad_weight": "Вес критерия «{name}» должен быть числом, не меньшим 0.",
 "st_dup_name": "Два критерия не могут называться одинаково.",
 "st_bad_scale": "Максимальная оценка должна быть от 2 до 10.",
 "font_title": "Шрифт", "font_family": "Шрифт", "font_size": "Размер",
 "font_preview": "Узнайте, что читать дальше — Aa Bb 123", "font_reset": "По умолчанию", "btn_apply": "Применить",
 "about_title": "О программе",
 "crit_1": "Интерес", "crit_2": "Личная польза", "crit_3": "Лёгкость чтения",
 "crit_4": "Социальный контекст", "crit_5": "Долгосрочная ценность / справочник",
 "ann_new_version": "Вышла новая версия: {latest} (у вас {current})", "ann_open_q": "Открыть ссылку в браузере?\n{url}",
 "ann_board_title": "Объявления",
 "ann_empty": "Сейчас объявлений нет.",
 "ann_type_update": "Обновление",
 "ann_type_new_app": "Новое приложение",
 "ann_type_recommended": "Рекомендуемое приложение",
 "ann_type_info": "Информация",
 "ann_open": "Открыть ссылку",
 "btn_close": "Закрыть",
 "about_hello": "Привет! Я Зеки. Я создаю Next Book, чтобы вы могли оценивать книги, которые хотите прочитать, по собственным критериям и видеть, что читать дальше.",
 "about_privacy": "Ваши данные остаются на вашем компьютере, учётная запись не нужна. Интернет используется только для поиска книг (Open Library) и объявлений.",
 "about_data": "Папка данных",
        "data_recovered": "Не удалось прочитать файл данных (возможно, он повреждён). Начат новый список; старый файл сохранён здесь:\n{path}",
},
"de": {
 "lbl_title": "Buchtitel", "lbl_author": "Autor",
 "btn_search": "Suchen", "btn_clear": "Leeren", "btn_manual": "+ Manuell hinzufügen",
 "results_title": " Suchergebnisse — Doppelklick zum Hinzufügen ",
 "btn_add_to_list": "Zur Liste hinzufügen", "list_title": " Meine Leseliste ",
 "btn_settings": "⚙ Bewertungseinstellungen", "btn_sort_score": "Nach Bewertung sortieren",
 "btn_delete": "Auswahl löschen", "btn_import": "Aus Excel importieren", "btn_export": "Nach Excel exportieren",
 "hint": "Doppelklick auf eine Zelle zum Bearbeiten · Enter = nach unten, Tab = nach rechts, Esc = abbrechen, leer = Bewertung löschen · "
         "Je kräftiger die Farbe, desto höher die Bewertung · Klick auf eine Überschrift sortiert · Rechtsklick auf ein Buch: Status, entfernen",
 "col_title": "Titel", "col_author": "Autor", "col_genre": "Genre", "col_pages": "Seiten",
 "col_year": "Jahr", "col_score": "Prioritätswert",
 "status_counts": "{total} Bücher · {scored} bewertet · {pending} warten auf Bewertung",
 "searching": "Suche läuft…", "enter_query": "Bitte einen Buchtitel oder Autor eingeben.",
 "no_results": "Keine Ergebnisse. Versuche eine andere Schreibweise oder „+ Manuell hinzufügen“.",
 "results_found": "{n} Ergebnisse gefunden. Doppelklick zum Hinzufügen.",
 "pick_result": "Bitte zuerst ein Ergebnis auswählen.",
 "err_network": "Open Library ist nicht erreichbar. Bitte Internetverbindung prüfen.",
 "err_search": "Suchfehler: {err}",
 "dup_book": "„{title}“ ist bereits in deiner Liste.", "ask_title": "Buchtitel:",
 "confirm_delete": "Die {n} ausgewählten Bücher aus der Liste löschen?",
 "bad_pages": "Die Seitenzahl muss eine Zahl sein.",
 "bad_score": "Ungültige Bewertung: Gib eine Zahl von 1 bis {max} ein.",
 "need_openpyxl": "Für Excel wird openpyxl benötigt.\n\nFühre in der Eingabeaufforderung aus:\n\npip install openpyxl",
 "pick_excel": "Excel-Datei auswählen",
 "confirm_import": "Deine aktuelle Liste wird durch die Bücher und Kriterien aus der Excel-Datei ersetzt.\n"
                   "(Sichere sie vorher ggf. mit „Nach Excel exportieren“.)\n\nFortfahren?",
 "import_failed": "Excel-Datei konnte nicht gelesen werden:\n{err}",
 "import_done": "{books} Bücher und {crit} Kriterien importiert.",
 "save_excel": "Als Excel speichern",
 "export_locked": "Datei konnte nicht geschrieben werden. Falls sie in Excel geöffnet ist, bitte schließen und erneut versuchen.",
 "export_failed": "Export fehlgeschlagen:\n{err}", "saved_to": "Gespeichert: {path}",
 "save_failed": "Speichern fehlgeschlagen:\n{err}",
 "xl_sheet_main": "Leseprioritäten", "xl_sheet_settings": "Einstellungen",
 "xl_weights_head": "Gewichtungen", "xl_weight": "Gewicht", "xl_total": "Summe",
 "xl_note": "Die Gewichte stehen im Verhältnis zueinander; die Summe muss nicht 1 ergeben.",
 "st_title": "Bewertungseinstellungen", "st_heading": "Lege deine eigenen Kriterien und Gewichte fest.",
 "st_desc": "Ein Kriterium mit höherem Gewicht beeinflusst die Gesamtwertung stärker. Die Gewichte stehen im "
            "Verhältnis zueinander; die Summe muss nicht 1 ergeben. Leere Bewertungen werden ignoriert.",
 "st_name": "Kriterium", "st_weight": "Gewicht", "st_add": "+ Kriterium hinzufügen",
 "st_scale": "Höchstwert:", "st_total": "Gesamtgewicht: {total}",
 "btn_cancel": "Abbrechen", "btn_save": "Speichern",
 "st_need_one": "Es muss mindestens ein Kriterium vorhanden sein.", "st_name_required": "Jedes Kriterium braucht einen Namen.",
 "st_bad_weight": "Das Gewicht für „{name}“ muss 0 oder größer sein.",
 "st_dup_name": "Zwei Kriterien dürfen nicht denselben Namen haben.",
 "st_bad_scale": "Der Höchstwert muss zwischen 2 und 10 liegen.",
 "font_title": "Schriftart", "font_family": "Schriftart", "font_size": "Größe",
 "font_preview": "Wissen, was als Nächstes dran ist — Aa Bb 123", "font_reset": "Standard", "btn_apply": "Anwenden",
 "about_title": "Info",
 "crit_1": "Interesse", "crit_2": "Persönlicher Nutzen", "crit_3": "Leichte Lesbarkeit",
 "crit_4": "Sozialer Kontext", "crit_5": "Langfristiger / Referenzwert",
 "ann_new_version": "Neue Version verfügbar: {latest} (installiert: {current})", "ann_open_q": "Link im Browser öffnen?\n{url}",
 "ann_board_title": "Ankündigungen",
 "ann_empty": "Derzeit keine Ankündigungen.",
 "ann_type_update": "Update",
 "ann_type_new_app": "Neue App",
 "ann_type_recommended": "Empfohlene App",
 "ann_type_info": "Info",
 "ann_open": "Link öffnen",
 "btn_close": "Schließen",
 "about_hello": "Hallo! Ich bin Zeki. Ich entwickle Next Book, damit du die Bücher, die du lesen möchtest, nach eigenen Kriterien bewerten kannst und siehst, was als Nächstes dran ist.",
 "about_privacy": "Deine Daten bleiben auf deinem Computer, ein Konto ist nicht nötig. Das Internet wird nur für die Buchsuche (Open Library) und für Ankündigungen genutzt.",
 "about_data": "Datenordner",
        "data_recovered": "Deine Datendatei konnte nicht gelesen werden (möglicherweise beschädigt). Es wurde eine neue Liste angelegt; die alte Datei wurde hier gesichert:\n{path}",
},
"fr": {
 "lbl_title": "Titre du livre", "lbl_author": "Auteur",
 "btn_search": "Rechercher", "btn_clear": "Effacer", "btn_manual": "+ Ajouter manuellement",
 "results_title": " Résultats — double-clic pour ajouter ",
 "btn_add_to_list": "Ajouter à la liste", "list_title": " Ma liste de lecture ",
 "btn_settings": "⚙ Réglages de notation", "btn_sort_score": "Trier par score",
 "btn_delete": "Supprimer la sélection", "btn_import": "Importer depuis Excel", "btn_export": "Exporter vers Excel",
 "hint": "Double-clic sur une cellule pour modifier · Entrée = bas, Tab = droite, Échap = annuler, vide = effacer la note · "
         "Plus la couleur est intense, plus le score est élevé · Cliquez sur un en-tête pour trier · Clic droit sur un livre : statut, retirer",
 "col_title": "Titre", "col_author": "Auteur", "col_genre": "Genre", "col_pages": "Pages",
 "col_year": "Année", "col_score": "Score de priorité",
 "status_counts": "{total} livres · {scored} notés · {pending} en attente de note",
 "searching": "Recherche…", "enter_query": "Saisissez un titre de livre ou un auteur.",
 "no_results": "Aucun résultat. Essayez une autre orthographe ou « + Ajouter manuellement ».",
 "results_found": "{n} résultats. Double-cliquez pour ajouter.",
 "pick_result": "Sélectionnez d'abord un résultat.",
 "err_network": "Impossible de joindre Open Library. Vérifiez votre connexion internet.",
 "err_search": "Erreur de recherche : {err}",
 "dup_book": "« {title} » est déjà dans votre liste.", "ask_title": "Titre du livre :",
 "confirm_delete": "Supprimer les {n} livres sélectionnés de la liste ?",
 "bad_pages": "Le nombre de pages doit être un nombre.",
 "bad_score": "Note invalide : saisissez un nombre entre 1 et {max}.",
 "need_openpyxl": "La prise en charge d'Excel nécessite openpyxl.\n\nExécutez dans l'invite de commandes :\n\npip install openpyxl",
 "pick_excel": "Choisir un fichier Excel",
 "confirm_import": "Votre liste actuelle sera remplacée par les livres et critères du fichier Excel.\n"
                   "(Vous pouvez d'abord faire une sauvegarde avec « Exporter vers Excel ».)\n\nContinuer ?",
 "import_failed": "Impossible de lire le fichier Excel :\n{err}",
 "import_done": "{books} livres et {crit} critères importés.",
 "save_excel": "Enregistrer au format Excel",
 "export_locked": "Impossible d'écrire le fichier. S'il est ouvert dans Excel, fermez-le et réessayez.",
 "export_failed": "Échec de l'export :\n{err}", "saved_to": "Enregistré : {path}",
 "save_failed": "Échec de l'enregistrement :\n{err}",
 "xl_sheet_main": "Priorités de lecture", "xl_sheet_settings": "Réglages",
 "xl_weights_head": "Pondérations", "xl_weight": "Poids", "xl_total": "Total",
 "xl_note": "Les poids sont relatifs ; leur somme n'a pas besoin d'être égale à 1.",
 "st_title": "Réglages de notation", "st_heading": "Définissez vos propres critères et pondérations.",
 "st_desc": "Un critère à fort poids influence davantage le score final. Les poids sont relatifs et leur somme "
            "n'a pas besoin d'être égale à 1. Les notes vides sont ignorées.",
 "st_name": "Critère", "st_weight": "Poids", "st_add": "+ Ajouter un critère",
 "st_scale": "Note maximale :", "st_total": "Poids total : {total}",
 "btn_cancel": "Annuler", "btn_save": "Enregistrer",
 "st_need_one": "Il faut au moins un critère.", "st_name_required": "Chaque critère doit avoir un nom.",
 "st_bad_weight": "Le poids de « {name} » doit être supérieur ou égal à 0.",
 "st_dup_name": "Deux critères ne peuvent pas avoir le même nom.",
 "st_bad_scale": "La note maximale doit être comprise entre 2 et 10.",
 "font_title": "Police", "font_family": "Police", "font_size": "Taille",
 "font_preview": "Savoir quoi lire ensuite — Aa Bb 123", "font_reset": "Par défaut", "btn_apply": "Appliquer",
 "about_title": "À propos",
 "crit_1": "Intérêt", "crit_2": "Apport personnel", "crit_3": "Facilité de lecture",
 "crit_4": "Contexte social", "crit_5": "Valeur à long terme / de référence",
 "ann_new_version": "Nouvelle version disponible : {latest} (vous avez {current})", "ann_open_q": "Ouvrir ce lien dans le navigateur ?\n{url}",
 "ann_board_title": "Annonces",
 "ann_empty": "Aucune annonce pour le moment.",
 "ann_type_update": "Mise à jour",
 "ann_type_new_app": "Nouvelle application",
 "ann_type_recommended": "Application recommandée",
 "ann_type_info": "Info",
 "ann_open": "Ouvrir le lien",
 "btn_close": "Fermer",
 "about_hello": "Bonjour ! Moi, c'est Zeki. Je développe Next Book pour vous aider à noter les livres que vous voulez lire selon vos propres critères et à savoir quoi lire ensuite.",
 "about_privacy": "Vos données restent sur votre ordinateur, aucun compte n'est nécessaire. Internet n'est utilisé que pour la recherche de livres (Open Library) et les annonces.",
 "about_data": "Dossier de données",
        "data_recovered": "Votre fichier de données est illisible (peut-être corrompu). Une nouvelle liste a été créée ; l'ancien fichier a été sauvegardé ici :\n{path}",
},
"zh": {
 "lbl_title": "书名", "lbl_author": "作者",
 "btn_search": "搜索", "btn_clear": "清除", "btn_manual": "+ 手动添加",
 "results_title": " 搜索结果 — 双击添加 ",
 "btn_add_to_list": "加入列表", "list_title": " 我的阅读列表 ",
 "btn_settings": "⚙ 评分设置", "btn_sort_score": "按得分排序",
 "btn_delete": "删除所选", "btn_import": "从 Excel 导入", "btn_export": "导出到 Excel",
 "hint": "双击单元格进行编辑 · Enter = 向下，Tab = 向右，Esc = 取消，留空 = 清除分数 · "
         "颜色越深，得分越高 · 点击表头可排序 · 右键点击书籍：状态、删除",
 "col_title": "书名", "col_author": "作者", "col_genre": "类型", "col_pages": "页数",
 "col_year": "年份", "col_score": "优先级得分",
 "status_counts": "共 {total} 本 · 已评分 {scored} 本 · 待评分 {pending} 本",
 "searching": "正在搜索…", "enter_query": "请输入书名或作者。",
 "no_results": "没有找到结果。请换个写法，或使用“+ 手动添加”。",
 "results_found": "找到 {n} 条结果，双击即可添加。",
 "pick_result": "请先选择一条结果。",
 "err_network": "无法连接 Open Library，请检查网络连接。",
 "err_search": "搜索出错：{err}",
 "dup_book": "《{title}》已在你的列表中。", "ask_title": "书名：",
 "confirm_delete": "要从列表中删除所选的 {n} 本书吗？",
 "bad_pages": "页数必须是数字。",
 "bad_score": "分数无效：请输入 1 到 {max} 之间的数字。",
 "need_openpyxl": "使用 Excel 功能需要 openpyxl。\n\n请在命令行中运行：\n\npip install openpyxl",
 "pick_excel": "选择 Excel 文件",
 "confirm_import": "当前列表将被 Excel 文件中的图书和评分标准替换。\n（建议先用“导出到 Excel”备份。）\n\n是否继续？",
 "import_failed": "无法读取 Excel 文件：\n{err}",
 "import_done": "已导入 {books} 本书和 {crit} 个评分标准。",
 "save_excel": "另存为 Excel",
 "export_locked": "无法写入文件。如果它在 Excel 中打开，请先关闭再重试。",
 "export_failed": "导出失败：\n{err}", "saved_to": "已保存：{path}",
 "save_failed": "保存失败：\n{err}",
 "xl_sheet_main": "阅读优先级", "xl_sheet_settings": "设置",
 "xl_weights_head": "权重", "xl_weight": "权重", "xl_total": "合计",
 "xl_note": "权重只看相对大小，总和不必等于 1。",
 "st_title": "评分设置", "st_heading": "自定义你的评分标准和权重。",
 "st_desc": "权重越高的标准，对最终得分影响越大。权重只看相对大小，总和不必等于 1。留空的分数不参与计算。",
 "st_name": "评分标准", "st_weight": "权重", "st_add": "+ 添加标准",
 "st_scale": "最高分：", "st_total": "权重合计：{total}",
 "btn_cancel": "取消", "btn_save": "保存",
 "st_need_one": "至少需要保留一个评分标准。", "st_name_required": "每个评分标准都需要名称。",
 "st_bad_weight": "“{name}”的权重必须是大于或等于 0 的数字。",
 "st_dup_name": "两个评分标准不能同名。",
 "st_bad_scale": "最高分必须在 2 到 10 之间。",
 "font_title": "字体", "font_family": "字体", "font_size": "字号",
 "font_preview": "知道下一本读什么 — Aa Bb 123", "font_reset": "默认", "btn_apply": "应用",
 "about_title": "关于",
 "crit_1": "兴趣程度", "crit_2": "个人收获", "crit_3": "阅读难易",
 "crit_4": "社交相关", "crit_5": "长期 / 参考价值",
 "ann_new_version": "新版本已发布：{latest}（当前：{current}）", "ann_open_q": "在浏览器中打开此链接？\n{url}",
 "ann_board_title": "公告",
 "ann_empty": "目前没有公告。",
 "ann_type_update": "更新",
 "ann_type_new_app": "新应用",
 "ann_type_recommended": "推荐应用",
 "ann_type_info": "信息",
 "ann_open": "打开链接",
 "btn_close": "关闭",
 "about_hello": "你好！我是 Zeki。我开发 Next Book，是为了让你按自己的标准给想读的书打分，并清楚下一本该读什么。",
 "about_privacy": "你的数据只保存在你的电脑上，无需账号。网络仅用于图书搜索（Open Library）和公告。",
 "about_data": "数据文件夹",
        "data_recovered": "无法读取你的数据文件（可能已损坏）。已创建新的列表，旧文件已备份到：\n{path}",
},
}

_TAB_STRINGS = {
"tr": {"tab_new": "Yeni liste", "tab_rename": "Yeniden adlandır", "tab_duplicate": "Çoğalt",
       "tab_close": "Listeyi sil", "ask_list_name": "Liste adı:",
       "confirm_delete_list": "“{name}” listesi ve içindeki {n} kitap silinsin mi?",
       "tab_limit": "En fazla {n} liste açılabilir.",
       "import_done_tab": "“{name}” listesi eklendi: {books} kitap, {crit} ölçüt."},
"en": {"tab_new": "New list", "tab_rename": "Rename", "tab_duplicate": "Duplicate",
       "tab_close": "Delete list", "ask_list_name": "List name:",
       "confirm_delete_list": "Delete the list “{name}” and its {n} books?",
       "tab_limit": "You can have at most {n} lists.",
       "import_done_tab": "Added list “{name}”: {books} books, {crit} criteria."},
"ru": {"tab_new": "Новый список", "tab_rename": "Переименовать", "tab_duplicate": "Дублировать",
       "tab_close": "Удалить список", "ask_list_name": "Название списка:",
       "confirm_delete_list": "Удалить список «{name}» и книги в нём ({n})?",
       "tab_limit": "Можно создать не более {n} списков.",
       "import_done_tab": "Добавлен список «{name}»: книг — {books}, критериев — {crit}."},
"de": {"tab_new": "Neue Liste", "tab_rename": "Umbenennen", "tab_duplicate": "Duplizieren",
       "tab_close": "Liste löschen", "ask_list_name": "Listenname:",
       "confirm_delete_list": "Liste „{name}“ mit {n} Büchern löschen?",
       "tab_limit": "Es sind höchstens {n} Listen möglich.",
       "import_done_tab": "Liste „{name}“ hinzugefügt: {books} Bücher, {crit} Kriterien."},
"fr": {"tab_new": "Nouvelle liste", "tab_rename": "Renommer", "tab_duplicate": "Dupliquer",
       "tab_close": "Supprimer la liste", "ask_list_name": "Nom de la liste :",
       "confirm_delete_list": "Supprimer la liste « {name} » et ses {n} livres ?",
       "tab_limit": "Vous pouvez avoir {n} listes au maximum.",
       "import_done_tab": "Liste « {name} » ajoutée : {books} livres, {crit} critères."},
"zh": {"tab_new": "新列表", "tab_rename": "重命名", "tab_duplicate": "复制",
       "tab_close": "删除列表", "ask_list_name": "列表名称：",
       "confirm_delete_list": "要删除列表“{name}”及其中的 {n} 本书吗？",
       "tab_limit": "最多只能有 {n} 个列表。",
       "import_done_tab": "已添加列表“{name}”：{books} 本书，{crit} 个评分标准。"},
}
for _code, _d in _TAB_STRINGS.items():
    I18N[_code].update(_d)
    I18N[_code]["tab_default"] = I18N[_code]["list_title"].strip()   # ilk listenin varsayılan adı

_ROW_STRINGS = {
"tr": {"st_todo": "Okunacak", "st_reading": "Okuyorum", "st_done": "Okudum",
       "col_status": "Durum", "row_delete": "Listeden sil"},
"en": {"st_todo": "To read", "st_reading": "Reading", "st_done": "Read",
       "col_status": "Status", "row_delete": "Remove from list"},
"ru": {"st_todo": "Хочу прочитать", "st_reading": "Читаю", "st_done": "Прочитано",
       "col_status": "Статус", "row_delete": "Удалить из списка"},
"de": {"st_todo": "Will ich lesen", "st_reading": "Lese ich gerade", "st_done": "Gelesen",
       "col_status": "Status", "row_delete": "Aus der Liste entfernen"},
"fr": {"st_todo": "À lire", "st_reading": "En cours", "st_done": "Lu",
       "col_status": "Statut", "row_delete": "Retirer de la liste"},
"zh": {"st_todo": "想读", "st_reading": "在读", "st_done": "已读",
       "col_status": "状态", "row_delete": "从列表中删除"},
}
for _code, _d in _ROW_STRINGS.items():
    I18N[_code].update(_d)

_NEW_STRINGS = {
"tr": {"st_paused": "Yarım bıraktım", "btn_grid": "Izgara",
       "hint_blank": "Alttaki boş satıra yazarak kitap ekle"},
"en": {"st_paused": "Paused", "btn_grid": "Grid",
       "hint_blank": "Type in the empty bottom row to add a book"},
"ru": {"st_paused": "Приостановлено", "btn_grid": "Сетка",
       "hint_blank": "Чтобы добавить книгу, пишите в пустой нижней строке"},
"de": {"st_paused": "Pausiert", "btn_grid": "Gitter",
       "hint_blank": "Neue Bücher in der leeren Zeile unten eintragen"},
"fr": {"st_paused": "En pause", "btn_grid": "Grille",
       "hint_blank": "Écrivez dans la ligne vide du bas pour ajouter un livre"},
"zh": {"st_paused": "已暂停", "btn_grid": "网格",
       "hint_blank": "在底部空行中输入即可添加书籍"},
}
for _code, _d in _NEW_STRINGS.items():
    I18N[_code].update(_d)

_LANG = ["en"]


def set_language(code):
    if code in I18N:
        _LANG[0] = code


def current_language():
    return _LANG[0]


def T(key, **kw):
    text = I18N.get(_LANG[0], {}).get(key)
    if text is None:
        text = I18N["en"].get(key, key)
    return text.format(**kw) if kw else text


def _aliases(key, extra=()):
    s = {v[key].lower() for v in I18N.values()}
    s.update(extra)
    return s


_AL_TITLE = _aliases("col_title", ("kitap adı", "kitap", "başlık", "title", "book title"))
_AL_AUTHOR = _aliases("col_author", ("yazar", "author"))
_AL_GENRE = _aliases("col_genre", ("tür", "kategori", "genre", "category"))
_AL_PAGES = _aliases("col_pages", ("sayfa sayısı", "sayfa", "pages", "page count"))
_AL_STATUS = _aliases("col_status", ("durum", "status"))
_AL_SCORE = _aliases("col_score", ("öncelik skoru", "skor", "priority score", "score"))


# ============================================================================
# TEMALAR
# ============================================================================

THEMES = {
    # ---- aydınlık temalar
    "cream": dict(   # kütüphane katalog çekmecesi (varsayılan)
        dark=False, swatch="#1F3A5F",
        bg="#F5F0E6", text="#2B2B2B", muted="#6B665A", hint="#8A8472", border="#D5CBB2",
        header_bg="#1F3A5F", header_fg="#FFFFFF", header_active="#2B4F80", title_fg="#1F3A5F",
        accent="#C8962E", accent_active="#DDAA3C", accent_fg="#1B1B1B",
        button="#E6DDC8", button_active="#DCD0B3", button_disabled="#EEE8DA",
        field="#FFFFFF", field_fg="#2B2B2B", tree_bg="#FFFFFF", unscored_fg="#9AA3B2",
        scored_fg="#143020", select_bg="#E7B94F", select_fg="#000000", label_frame="#6B5A33",
        ramp=["#EAF2E6", "#D6E7D0", "#C0DAB8", "#A8CC9F", "#8FBC85"]),
    "moss": dict(
        dark=False, swatch="#35553B",
        bg="#EDF1E8", text="#25301F", muted="#5E6B55", hint="#7F8B76", border="#C9D3BD",
        header_bg="#35553B", header_fg="#FFFFFF", header_active="#46704E", title_fg="#35553B",
        accent="#C9A227", accent_active="#DDB63A", accent_fg="#1F1A05",
        button="#DDE6D2", button_active="#CEDAC0", button_disabled="#E7EDDF",
        field="#FFFFFF", field_fg="#25301F", tree_bg="#FFFFFF", unscored_fg="#9BA793",
        scored_fg="#25330F", select_bg="#E7B94F", select_fg="#000000", label_frame="#4B6B3F",
        ramp=["#EEF3E2", "#E0EBC8", "#CFE0A8", "#BBD28A", "#A4C26B"]),
    "ocean": dict(
        dark=False, swatch="#14506B",
        bg="#E9F1F6", text="#16303F", muted="#587080", hint="#7A93A3", border="#BCD3E0",
        header_bg="#14506B", header_fg="#FFFFFF", header_active="#1E6A8A", title_fg="#14506B",
        accent="#E08A2E", accent_active="#EE9C45", accent_fg="#1B1B1B",
        button="#D7E6EF", button_active="#C7DCE9", button_disabled="#E2EDF3",
        field="#FFFFFF", field_fg="#16303F", tree_bg="#FFFFFF", unscored_fg="#9AA9B5",
        scored_fg="#0E2F44", select_bg="#F2C14E", select_fg="#000000", label_frame="#2D6580",
        ramp=["#E6F1F9", "#D0E5F3", "#B8D8EE", "#9CC8E6", "#7DB4DA"]),
    # ---- karanlık temalar
    "night": dict(
        dark=True, swatch="#3B4B6B",
        bg="#171C26", text="#E4E7EC", muted="#98A2B3", hint="#7B8598", border="#303A4C",
        header_bg="#0E131B", header_fg="#E4E7EC", header_active="#1D2736", title_fg="#F2C14E",
        accent="#E0A93B", accent_active="#F0BC55", accent_fg="#17120A",
        button="#263043", button_active="#31405A", button_disabled="#1F2736",
        field="#202838", field_fg="#E4E7EC", tree_bg="#1D2431", unscored_fg="#66728A",
        scored_fg="#EAF5EA", select_bg="#B9852A", select_fg="#FFFFFF", label_frame="#D6B25E",
        ramp=["#24362F", "#2A4637", "#315640", "#38674A", "#407955"]),
    "plum": dict(
        dark=True, swatch="#6A4FB0",
        bg="#1D1828", text="#E9E4F3", muted="#A79FC0", hint="#8B83A6", border="#383050",
        header_bg="#120E1B", header_fg="#E9E4F3", header_active="#262037", title_fg="#C8B3F5",
        accent="#A98BEA", accent_active="#BBA2F2", accent_fg="#150F26",
        button="#2B2440", button_active="#382F55", button_disabled="#231D33",
        field="#251F36", field_fg="#E9E4F3", tree_bg="#211B30", unscored_fg="#6F6890",
        scored_fg="#F4F0FC", select_bg="#7B5CC4", select_fg="#FFFFFF", label_frame="#C0A9F0",
        ramp=["#2D2545", "#382C58", "#44366C", "#514180", "#5F4E96"]),
    "ember": dict(
        dark=True, swatch="#B8662A",
        bg="#221A15", text="#F0E6DC", muted="#B3A092", hint="#8F7D70", border="#43342B",
        header_bg="#150F0B", header_fg="#F0E6DC", header_active="#2E221A", title_fg="#F0B070",
        accent="#E08A3C", accent_active="#EE9D55", accent_fg="#1C1008",
        button="#33261E", button_active="#443328", button_disabled="#2A2019",
        field="#2B2018", field_fg="#F0E6DC", tree_bg="#271D17", unscored_fg="#7C6B5E",
        scored_fg="#FFF4E8", select_bg="#4F86A8", select_fg="#FFFFFF", label_frame="#E0A468",
        ramp=["#3A2A1C", "#4A3320", "#5B3D24", "#6D4929", "#80552E"]),
}
DEFAULT_THEME = "cream"


def resource_path(name):
    """PyInstaller paketinde de, düz çalıştırmada da dosyayı bul."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, name)


def score_bucket(score, smax, steps=5):
    """Skoru 0..steps-1 arası bir renk kademesine çevir."""
    if smax <= 0:
        return steps - 1
    t = max(0.0, min(1.0, score / smax))   # 10 üzerinden: 0–2, 2–4, 4–6, 6–8, 8–10
    return min(steps - 1, int(t * steps))


def short(text, n):
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def fit_text(font, text, px):
    """Metni px piksele sığacak şekilde sonuna … koyarak kısalt."""
    if px <= 10 or font.measure(text) <= px:
        return text
    lo, hi = 0, len(text)
    while lo < hi:
        mid = (lo + hi + 1) // 2
        if font.measure(text[:mid].rstrip() + "…") <= px:
            lo = mid
        else:
            hi = mid - 1
    return text[:lo].rstrip() + "…"


def _mix(h1, h2, t):
    """İki #RRGGBB rengi karıştır (t=0 → h1, t=1 → h2)."""
    a, b = _rgb(h1), _rgb(h2)
    return "#%02X%02X%02X" % tuple(round(x + (y - x) * t) for x, y in zip(a, b))



# ============================================================================
# BAYRAKLAR (resim dosyası yok: piksel piksel çizilir, sonra PhotoImage olur)
# ============================================================================

_RGB_CACHE = {}


def _rgb(h):
    c = _RGB_CACHE.get(h)
    if c is None:
        c = (int(h[1:3], 16), int(h[3:5], 16), int(h[5:7], 16))
        _RGB_CACHE[h] = c
    return c


# ızgara çizgisi rengi (_rgb tanımından sonra hesaplanır)
for _t in THEMES.values():   # ızgara çizgisi: satır zeminine göre çok hafif bir ton farkı
    _t["grid"] = _mix(_t["tree_bg"], "#FFFFFF" if _t["dark"] else _t["text"], 0.16 if _t["dark"] else 0.14)


def _star_points(cx, cy, R, rot_deg):
    pts = []
    for i in range(10):
        a = math.radians(rot_deg + i * 36)
        rad = R if i % 2 == 0 else R * 0.382
        pts.append((cx + rad * math.cos(a), cy - rad * math.sin(a)))
    return pts


def _make_star(cx, cy, R, rot_deg):
    return (cx, cy, R * R, _star_points(cx, cy, R, rot_deg))


def _in_star(x, y, star):
    cx, cy, r2, pts = star
    if (x - cx) ** 2 + (y - cy) ** 2 > r2:
        return False
    inside = False
    j = len(pts) - 1
    for i in range(len(pts)):
        xi, yi = pts[i]
        xj, yj = pts[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


_TR_STAR = _make_star(16.0, 10.0, 2.5, 180)
_CN_STARS = [_make_star(5, 5, 3, 90)] + [
    _make_star(cx, cy, 1, math.degrees(math.atan2(cy - 5, 5 - cx)))
    for cx, cy in ((10, 2), (12, 4), (12, 7), (10, 9))
]


def _flag_color(code, u, v):
    if code == "de":
        return "#000000" if v < 1 / 3 else ("#DD0000" if v < 2 / 3 else "#FFCE00")
    if code == "fr":
        return "#0055A4" if u < 1 / 3 else ("#FFFFFF" if u < 2 / 3 else "#EF4135")
    if code == "ru":
        return "#FFFFFF" if v < 1 / 3 else ("#0039A6" if v < 2 / 3 else "#D52B1E")
    if code == "tr":
        x, y = u * 30, v * 20
        if (x - 9.5) ** 2 + (y - 10) ** 2 <= 25 and (x - 10.75) ** 2 + (y - 10) ** 2 > 16:
            return "#FFFFFF"
        if _in_star(x, y, _TR_STAR):
            return "#FFFFFF"
        return "#E30A17"
    if code == "cn" or code == "zh":
        x, y = u * 30, v * 20
        for st in _CN_STARS:
            if _in_star(x, y, st):
                return "#FFDE00"
        return "#DE2910"
    if code == "en" or code == "gb":
        x, y = u * 60, v * 40
        n = 72.111
        d1 = abs(40 * x - 60 * y) / n
        d2 = abs(40 * x + 60 * y - 2400) / n
        cx, cy = abs(x - 30), abs(y - 20)
        if cx <= 2.4 or cy <= 2.4:
            return "#C8102E"
        if cx <= 4 or cy <= 4:
            return "#FFFFFF"
        if d1 <= 1.2 or d2 <= 1.2:
            return "#C8102E"
        if d1 <= 3 or d2 <= 3:
            return "#FFFFFF"
        return "#012169"
    return "#888888"


def flag_rows(code, w, h):
    """Bayrağı w×h piksellik, kenarları yumuşatılmış hex renk satırları olarak üretir."""
    ss = 3 if w * h <= 1200 else 2
    n = ss * ss
    rows = []
    for py in range(h):
        row = []
        for px in range(w):
            r = g = b = 0
            for sy in range(ss):
                for sx in range(ss):
                    c = _rgb(_flag_color(code, (px + (sx + .5) / ss) / w, (py + (sy + .5) / ss) / h))
                    r += c[0]
                    g += c[1]
                    b += c[2]
            r, g, b = r // n, g // n, b // n
            if px in (0, w - 1) or py in (0, h - 1):      # ince koyu çerçeve
                r, g, b = int(r * .6 + 90 * .4), int(g * .6 + 90 * .4), int(b * .6 + 90 * .4)
            row.append("#%02x%02x%02x" % (r, g, b))
        rows.append(row)
    return rows


def make_flag_image(code, w, h):
    img = tk.PhotoImage(width=w, height=h)
    rows = flag_rows(code, w, h)
    img.put(" ".join("{" + " ".join(r) + "}" for r in rows))
    return img


# ============================================================================
# DUYURU SİMGELERİ (yazı tipi karakteri yerine çizilir: her bilgisayarda aynı ve tam ortalı)
# ============================================================================

ANN_TYPES = ("update", "new_app", "recommended", "info")


def ann_icon_colors(t, kind):
    """(daire rengi, simge rengi) — temaya uyar."""
    return {"update": (t["accent"], t["accent_fg"]),
            "new_app": (t["header_bg"], t["header_fg"]),
            "recommended": (t["label_frame"], t["bg"])}.get(kind, (t["muted"], t["bg"]))


def _in_poly(x, y, pts):
    inside, j = False, len(pts) - 1
    for i in range(len(pts)):
        xi, yi = pts[i]
        xj, yj = pts[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def ann_icon_rows(kind, size, fill, glyph, bg):
    """Yuvarlak rozet: info = i, update = yukarı ok, new_app = yıldız, recommended = kalp.
    4x4 alt örnekleme ile kenarlar yumuşatılır; satırlar '#rrggbb' dizileridir."""
    S = float(size)
    c = S / 2.0
    R = S / 2.0 - 0.5
    star = _star_points(c, c + S * 0.015, S * 0.30, 90)
    hs = S * 0.2368

    def in_glyph(x, y):
        if kind == "info":
            return ((x - c) ** 2 + (y - S * 0.31) ** 2 <= (S * 0.065) ** 2
                    or (abs(x - c) <= S * 0.065 and S * 0.45 <= y <= S * 0.76))
        if kind == "update":
            if abs(x - c) <= S * 0.07 and S * 0.48 <= y <= S * 0.76:
                return True
            return S * 0.24 <= y <= S * 0.50 and abs(x - c) <= (y - S * 0.24) * 0.21 / 0.26
        if kind == "new_app":
            return _in_poly(x, y, star)
        hx, hy = (x - c) / hs, -(y - c) / hs + 0.125      # recommended: kalp eğrisi
        return (hx * hx + hy * hy - 1) ** 3 - hx * hx * hy ** 3 <= 0

    cf, cg, cb = _rgb(fill), _rgb(glyph), _rgb(bg)
    n, rows = 4, []
    for py in range(size):
        row = []
        for px in range(size):
            r = g = b = 0
            for sy in range(n):
                for sx in range(n):
                    x, y = px + (sx + 0.5) / n, py + (sy + 0.5) / n
                    if (x - c) ** 2 + (y - c) ** 2 > R * R:
                        col = cb
                    elif in_glyph(x, y):
                        col = cg
                    else:
                        col = cf
                    r += col[0]
                    g += col[1]
                    b += col[2]
            k = n * n
            row.append("#%02x%02x%02x" % (round(r / k), round(g / k), round(b / k)))
        rows.append(row)
    return rows


def make_ann_icon(kind, size, fill, glyph, bg):
    img = tk.PhotoImage(width=size, height=size)
    rows = ann_icon_rows(kind, size, fill, glyph, bg)
    img.put(" ".join("{" + " ".join(r) + "}" for r in rows))
    return img


# ============================================================================
# DURUM ROZETLERİ (okuyorum = oynat, okudum = onay) — satır renginden bağımsız, saydam zeminli
# ============================================================================

def _seg_dist(x, y, a, b):
    ax, ay = a
    bx, by = b
    dx, dy = bx - ax, by - ay
    t = max(0.0, min(1.0, ((x - ax) * dx + (y - ay) * dy) / (dx * dx + dy * dy)))
    return math.hypot(x - (ax + t * dx), y - (ay + t * dy))


def _png_chunk(tag, data):
    body = tag + data
    return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)


def status_icon_png(kind, size, fill, glyph, ring=None, width=None):
    """Yuvarlak rozet (kind: 'reading' = oynat üçgeni, 'paused' = duraklat çubukları, 'done' = onay işareti) → RGBA PNG baytları.
    4x4 alt örnekleme: kenar yumuşak, zemin saydam; hangi renkli satırın üstünde olursa olsun uyar."""
    S = float(size)
    W = int(width or size)
    ox = (W - size) / 2.0   # rozet, görselin yatay ortasında durur; kenar boşlukları saydamdır
    c, R = S / 2.0, S / 2.0 - 0.5
    rw = max(1.0, S * 0.07) if ring else 0.0
    tri = [(S * 0.39, S * 0.27), (S * 0.39, S * 0.73), (S * 0.745, S * 0.50)]
    chk = [((S * 0.27, S * 0.52), (S * 0.43, S * 0.68)), ((S * 0.43, S * 0.68), (S * 0.74, S * 0.34))]
    hw = S * 0.075

    def in_glyph(x, y):
        if kind == "reading":
            return _in_poly(x, y, tri)
        if kind == "paused":   # ⏸ iki dikey çubuk
            return (S * 0.32 <= x <= S * 0.45 or S * 0.55 <= x <= S * 0.68) and S * 0.28 <= y <= S * 0.72
        return any(_seg_dist(x, y, a, b) <= hw for a, b in chk)

    cf, cg = _rgb(fill), _rgb(glyph)
    cr = _rgb(ring) if ring else cf
    n = 4
    raw = bytearray()
    for py in range(size):
        raw.append(0)   # PNG satır filtresi: yok
        for px in range(W):
            r = g = b = hit = 0
            for sy in range(n):
                for sx in range(n):
                    x, y = px - ox + (sx + 0.5) / n, py + (sy + 0.5) / n
                    d2 = (x - c) ** 2 + (y - c) ** 2
                    if d2 > R * R:
                        continue
                    col = cg if in_glyph(x, y) else (cr if d2 > (R - rw) ** 2 else cf)
                    r += col[0]
                    g += col[1]
                    b += col[2]
                    hit += 1
            if hit:
                raw += bytes((round(r / hit), round(g / hit), round(b / hit), round(255 * hit / (n * n))))
            else:
                raw += b"\x00\x00\x00\x00"
    return (b"\x89PNG\r\n\x1a\n"
            + _png_chunk(b"IHDR", struct.pack(">IIBBBBB", W, size, 8, 6, 0, 0, 0))
            + _png_chunk(b"IDAT", zlib.compress(bytes(raw), 9))
            + _png_chunk(b"IEND", b""))


def make_status_icon(kind, size, fill, glyph, ring=None, width=None):
    data = base64.b64encode(status_icon_png(kind, size, fill, glyph, ring, width)).decode("ascii")
    return tk.PhotoImage(data=data, format="png")


# ============================================================================
# OTOMATİK TAMAMLAMA (yazar / tür hücreleri)
# ============================================================================

def _fold(text):
    """Büyük/küçük harf ve i/İ/ı/I farkını yok say (uzunluk değişmez)."""
    return text.replace("İ", "i").replace("I", "i").replace("ı", "i").lower()


def suggest_values(books, key):
    """Listede daha önce yazılmış değerler: en sık kullanılan önce."""
    counts = {}
    for b in books:
        v = (b.get(key) or "").strip()
        if v:
            counts[v] = counts.get(v, 0) + 1
    return sorted(counts, key=lambda v: (-counts[v], _fold(v)))


def attach_autocomplete(entry, get_values):
    """Yazarken kalan kısmı seçili olarak tamamlar (Excel / tarayıcı gibi).
    Enter / Tab / → kabul eder, yazmaya devam etmek ya da Backspace önerinin yerine geçer."""
    def on_key(e):
        if not e.char or not e.char.isprintable() or (e.state & 0x4):   # Ctrl kombinasyonları hariç
            return
        text = entry.get()
        if not text.strip() or entry.index("insert") != len(text):
            return
        ft = _fold(text)
        for v in get_values():
            if len(v) > len(text) and _fold(v).startswith(ft):
                entry.delete(0, "end")
                entry.insert(0, v)
                entry.icursor(len(text))
                entry.selection_range(len(text), "end")
                return
    entry.bind("<KeyRelease>", on_key, add="+")
    return on_key


# ============================================================================
# YEREL ADLI VARSAYILAN ÖLÇÜTLER
# ============================================================================

BUILTIN_KEYS = [f"crit_{i}" for i in range(1, 6)]


def tag_builtin_criteria(criteria):
    """Adı herhangi bir dildeki varsayılan ölçüt adıyla aynı olanları 'builtin' diye işaretle.
    Böylece dil değişince adları otomatik yeni dile çevrilir; kullanıcının kendi
    yazdığı / yeniden adlandırdığı ölçütlere dokunulmaz."""
    lookup = {}
    for d in I18N.values():
        for k in BUILTIN_KEYS:
            lookup.setdefault(d[k].casefold(), k)
    for c in criteria:
        if c.get("builtin") not in BUILTIN_KEYS:
            c.pop("builtin", None)
            k = lookup.get(str(c.get("name", "")).strip().casefold())
            if k:
                c["builtin"] = k
    return criteria


def localize_criteria(criteria):
    """Varsayılan ölçütlerin adını geçerli dile çevir. Değişen bir şey olduysa True."""
    changed = False
    for c in criteria:
        k = c.get("builtin")
        if k in BUILTIN_KEYS and c.get("name") != T(k):
            c["name"] = T(k)
            changed = True
    return changed


# ============================================================================
# DUYURULAR (tüm uygulamaların ortak duyuru dosyası)
# ============================================================================

def parse_version(v):
    parts = re.findall(r"\d+", str(v))
    return tuple(int(x) for x in parts[:4]) or (0,)


def safe_url(u):
    u = str(u or "").strip()
    return u if u.lower().startswith("https://") and len(u) <= 500 else ""


def _loc(value, lang):
    """Çok dilli metinden kullanıcının dilini seç: önce lang, sonra en, sonra herhangi bir dil.
    Anahtarlar 'tr', 'zh-CN', 'zh_cn' gibi yazılmış olabilir; yalnızca ana dil koduna bakılır."""
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        by_lang = {}
        for k, v in value.items():
            if isinstance(v, str) and v.strip():
                by_lang.setdefault(str(k).strip().lower().replace("_", "-").split("-")[0], v)
        return by_lang.get(lang) or by_lang.get("en") or next(iter(by_lang.values()), "")
    return ""


class _HttpsOnlyRedirect(urllib.request.HTTPRedirectHandler):
    """https adresinden http adresine yönlendirmeyi reddet ("yalnızca HTTPS" sözü delinmesin)."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        if not newurl.lower().startswith("https://"):
            raise urllib.error.URLError("redirect to a non-https url was blocked")
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def fetch_feed(url):
    if not url.lower().startswith("https://"):
        raise ValueError("announcement url must be https")
    req = urllib.request.Request(url, headers={"User-Agent": f"NextBook/{APP_VERSION}",
                                               "Accept": "application/json"})
    with urllib.request.build_opener(_HttpsOnlyRedirect).open(req, timeout=6) as r:
        raw = r.read(65537)
    if len(raw) > 65536:
        raise ValueError("announcement feed too large")
    return json.loads(raw.decode("utf-8"))


def select_announcements(feed, app_id, version, lang, dismissed=(), today=None):
    """Bu uygulama + bu sürüm + bu dil için gösterilecek duyurular (öncelik sırasıyla).

    feed = {"apps": {"nextbook": {"latest": "1.3.0", "url": "https://..."}},
            "announcements": [{"id", "apps": ["nextbook"|"*"], "min_version", "max_version",
                               "starts", "expires", "priority", "type", "text": {"tr": "..", "en": ".."},
                               "url"}]}"""
    out = []
    if not isinstance(feed, dict):
        return out
    today = today or datetime.date.today()
    cur = parse_version(version)

    apps = feed.get("apps")
    info = apps.get(app_id) if isinstance(apps, dict) else None
    if isinstance(info, dict) and info.get("latest") and parse_version(info["latest"]) > cur:
        uid = f"update-{app_id}-{info['latest']}"
        if uid not in dismissed:
            out.append({"id": uid, "type": "update", "priority": 1000,
                        "text": T("ann_new_version", latest=info["latest"], current=version),
                        "url": safe_url(info.get("url")) or safe_url(REPO_URL + "/releases")})

    items = feed.get("announcements")
    for a in (items if isinstance(items, list) else []):
        try:
            if not isinstance(a, dict):
                continue
            aid = str(a.get("id") or "")[:80]
            if not aid or aid in dismissed:
                continue
            targets = a.get("apps", ["*"])
            targets = [targets] if isinstance(targets, str) else list(targets)
            if "*" not in targets and app_id not in targets:
                continue
            if a.get("min_version") and cur < parse_version(a["min_version"]):
                continue
            if a.get("max_version") and cur > parse_version(a["max_version"]):
                continue
            if a.get("starts") and today < datetime.date.fromisoformat(str(a["starts"])):
                continue
            if a.get("expires") and today > datetime.date.fromisoformat(str(a["expires"])):
                continue
            text = " ".join(_loc(a.get("text"), lang).split())[:600]
            if not text:
                continue
            kind = a.get("type") if a.get("type") in ANN_TYPES else "info"
            out.append({"id": aid, "type": kind, "text": text, "url": safe_url(a.get("url")),
                        "priority": int(a.get("priority", 0))})
        except (ValueError, TypeError):
            continue   # bozuk tek bir duyuru diğerlerini engellemesin
    out.sort(key=lambda x: -x["priority"])   # sort kararlıdır: eşit önceliklerde dosya sırası korunur
    return out


# ============================================================================
# VERİ MODELİ
# ============================================================================

def new_id():
    return uuid.uuid4().hex[:8]


def default_data(lang="tr"):
    names = I18N.get(lang, I18N["en"])
    return {
        "scale_max": DEFAULT_SCALE_MAX,
        "criteria": [{"id": new_id(), "name": names[f"crit_{i + 1}"], "weight": w,
                      "builtin": f"crit_{i + 1}"}
                     for i, w in enumerate(DEFAULT_WEIGHTS)],
        "books": [],
    }


def new_list(name, content=None, builtin=False):
    """Bir sekme = bir liste: kendi ölçütleri, puan üst sınırı ve kitapları var."""
    c = content if content is not None else default_data(current_language())
    lst = {"id": new_id(), "name": name, "scale_max": c["scale_max"],
           "criteria": c["criteria"], "books": c["books"]}
    if builtin:
        lst["name_builtin"] = True   # dil değişince adı da çevrilir; elle adlandırılınca kalkar
    return lst


def default_store(lang="tr"):
    names = I18N.get(lang, I18N["en"])
    lst = new_list(names["tab_default"], default_data(lang), builtin=True)
    return {"version": 2, "active": lst["id"], "lists": [lst]}


def active_list(store):
    return next((l for l in store["lists"] if l["id"] == store["active"]), store["lists"][0])


def new_book(title, author="", genre="", pages=None, key=""):
    return {
        "id": new_id(),
        "key": key,          # Open Library work key (tekrar eklemeyi önlemek için)
        "title": title,
        "author": author,
        "genre": genre,
        "pages": pages,
        "scores": {},        # {ölçüt_id: puan}
    }


def book_status(b):
    """'todo' (okunacak) · 'reading' (okuyorum) · 'done' (okudum). Alan yoksa okunacak sayılır."""
    s = b.get("status")
    return s if s in ("reading", "paused", "done") else "todo"


def compute_score(book, criteria):
    """Boş bırakılan ölçütler hesaba katılmaz:
    skor = Σ(puan × ağırlık) / Σ(dolu ölçütlerin ağırlığı)"""
    num = den = 0.0
    for c in criteria:
        v = book["scores"].get(c["id"])
        if v is None or v == "":
            continue
        num += float(v) * float(c["weight"])
        den += float(c["weight"])
    return round(num / den, 2) if den > 0 else 0.0


def is_scored(book, criteria):
    return any(book["scores"].get(c["id"]) not in (None, "") for c in criteria)


def fmt_num(v):
    if v is None or v == "":
        return ""
    v = float(v)
    return str(int(v)) if v == int(v) else f"{v:.2f}".rstrip("0").rstrip(".")


def parse_float(text):
    return float(text.strip().replace(",", "."))


LOAD_NOTICE = []   # veri dosyası okunamayıp yedeklendiyse yedeğin yolu buraya yazılır


def _validate_data(data, base):
    """Dosyadaki yapı beklenenden farklıysa ValueError fırlat (arayüz sonradan KeyError ile çökmesin)."""
    if not isinstance(data, dict):
        raise ValueError("root is not an object")
    data.setdefault("scale_max", base["scale_max"])
    data.setdefault("criteria", base["criteria"])
    data.setdefault("books", [])
    if not isinstance(data["criteria"], list) or not isinstance(data["books"], list):
        raise ValueError("criteria/books must be lists")
    smax = int(data["scale_max"])
    data["scale_max"] = smax if 2 <= smax <= 10 else base["scale_max"]
    if not data["criteria"]:
        data["criteria"] = base["criteria"]
    for c in data["criteria"]:
        if not isinstance(c, dict) or not isinstance(c.get("id"), str) or not isinstance(c.get("name"), str):
            raise ValueError("bad criterion")
        w = float(c.get("weight", 1.0))
        if not math.isfinite(w) or w < 0:
            raise ValueError("bad weight")
        c["weight"] = w
    for b in data["books"]:
        if not isinstance(b, dict) or not isinstance(b.get("id"), str) or not isinstance(b.get("title"), str):
            raise ValueError("bad book")
        b.setdefault("author", "")
        b.setdefault("genre", "")
        b.setdefault("pages", None)
        b.setdefault("key", "")
        if b.get("status") not in ("reading", "paused", "done"):
            b.pop("status", None)
        if not isinstance(b.get("scores"), dict):
            b["scores"] = {}
    return data


def _validate_store(data, lang):
    """Depo (sürüm 2) ya da eski tek listeli dosya (sürüm 1) gelir; (depo, taşındı_mı) döner."""
    if not isinstance(data, dict):
        raise ValueError("root is not an object")
    names = I18N.get(lang, I18N["en"])
    if "lists" not in data:   # eski dosya: tek liste → ilk sekme
        lst = _validate_data(data, default_data(lang))
        lst = new_list(names["tab_default"], lst, builtin=True)
        tag_builtin_criteria(lst["criteria"])
        return {"version": 2, "active": lst["id"], "lists": [lst]}, True
    if not isinstance(data["lists"], list):
        raise ValueError("lists must be a list")
    lists, seen = [], set()
    for l in data["lists"]:
        if not isinstance(l, dict):
            raise ValueError("bad list")
        _validate_data(l, default_data(lang))
        if not isinstance(l.get("id"), str) or l["id"] in seen:
            l["id"] = new_id()
        seen.add(l["id"])
        name = l.get("name")
        l["name"] = name.strip()[:60] if isinstance(name, str) and name.strip() else names["tab_new"]
        tag_builtin_criteria(l["criteria"])
        lists.append(l)
    if not lists:
        return default_store(lang), False
    data = {"version": 2, "active": data.get("active"), "lists": lists[:MAX_LISTS]}
    if data["active"] not in seen:
        data["active"] = lists[0]["id"]
    return data, False


def load_data(lang="tr"):
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            raw = json.load(f)
        store, migrated = _validate_store(raw, lang)
        if migrated:   # sürüm 1 dosyasına dokunmadan bir kez yedeğini al
            try:
                backup = DATA_FILE.with_name("next_book.v1-yedek.json")
                if not backup.exists():
                    backup.write_bytes(DATA_FILE.read_bytes())
            except OSError:
                pass
        return store
    except FileNotFoundError:
        return default_store(lang)
    except Exception:
        # Bozuk dosya: silme, zaman damgalı bir yedeğe taşı (eski yedeğin üstüne yazma) ve kullanıcıya haber ver.
        stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = DATA_FILE.with_name(f"next_book.bozuk-{stamp}.json")
        try:
            DATA_FILE.replace(backup)
        except Exception:
            backup = DATA_FILE
        LOAD_NOTICE.append(str(backup))
        return default_store(lang)


def _atomic_write(path, obj):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    tmp.replace(path)


def save_data(data):
    _atomic_write(DATA_FILE, data)


def load_settings():
    try:
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            d = json.load(f)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def save_settings(settings):
    try:
        _atomic_write(SETTINGS_FILE, settings)
    except OSError:
        pass


# ============================================================================
# OPEN LIBRARY ARAMA (anahtarsız) — sadece ad, yazar, tür, sayfa sayısı
# ============================================================================

_JUNK_SUBJECT_PARTS = (
    "accessible", "daisy", "in library", "lending", "overdrive", "open library",
    "large type", "reading level", "nyt:", "new york times", "internet archive",
    "protected", "bestseller", "award", "collection", "translations",
)
_PREFERRED_GENRES = (
    "science fiction", "fantasy", "fiction", "biography", "history", "philosophy",
    "psychology", "business", "economics", "self-help", "poetry", "mystery",
    "thriller", "romance", "horror", "science", "politics", "religion",
    "computers", "art", "health", "education", "mathematics", "sociology",
    "classic literature", "short stories", "essays", "drama", "humor",
)


def pick_genre(subjects):
    """Open Library konu listesi gürültülü; makul bir tür seç."""
    clean = []
    for s in subjects or []:
        t = str(s or "").strip()
        low = t.lower()
        if not t or len(t) > 30 or any(j in low for j in _JUNK_SUBJECT_PARTS):
            continue
        clean.append(t)
    for g in _PREFERRED_GENRES:
        for t in clean:
            if re.search(r"\b" + re.escape(g) + r"\b", t.lower()):
                return g.title()
    if clean:
        first = re.split(r"[,;]", clean[0])[0].strip()
        return first[:1].upper() + first[1:]
    return ""


def search_books(title, author, limit=40):
    params = {
        "limit": limit,
        "fields": "key,title,author_name,subject,number_of_pages_median,first_publish_year",
    }
    if title:
        params["title"] = title
    if author:
        params["author"] = author
    url = "https://openlibrary.org/search.json?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": f"NextBook/{APP_VERSION} (+{REPO_URL})"})
    with urllib.request.urlopen(req, timeout=15) as resp:
        payload = json.load(resp)
    return clean_results(parse_search_docs(payload.get("docs", [])), title)[:25]


def parse_search_docs(docs):
    results = []
    for d in docs:
        try:
            pages = int(d.get("number_of_pages_median") or 0) or None
        except (TypeError, ValueError, OverflowError):
            pages = None
        results.append({
            "key": str(d.get("key") or ""),
            "title": str(d.get("title") or "").strip(),
            "author": ", ".join(str(a) for a in (d.get("author_name") or [])[:3]),
            "genre": pick_genre(d.get("subject")),
            "pages": pages,
            "year": d.get("first_publish_year") or "",
        })
    return [r for r in results if r["title"]]


_NOISE_TITLE = re.compile(r"\[(?:collection|set|box|omnibus)|box(?:ed)? set|collection/set", re.I)


def _norm(text):
    return re.sub(r"[^\w]+", " ", (text or "").casefold()).strip()


def clean_results(results, query_title=""):
    """Takım/koleksiyon kayıtlarını ele, aynı kitabın tekrarlarını birleştir,
    aranan başlığa en yakın olanları öne al."""
    kept = [r for r in results if not _NOISE_TITLE.search(r["title"])]
    if not kept:
        kept = list(results)

    def info(r):
        return bool(r["pages"]) + bool(r["genre"]) + bool(r["year"]) + bool(r["author"])

    best, order = {}, []
    for r in kept:
        first_author = _norm((r["author"] or "").split(",")[0]).split()
        key = (_norm(r["title"]), first_author[-1] if first_author else "")
        if key not in best:
            best[key] = r
            order.append(key)
        elif info(r) > info(best[key]):
            best[key] = r
    merged = [best[k] for k in order]

    q = _norm(query_title)

    def rank(r):
        t = _norm(r["title"])
        if q and t == q:
            return 0
        if q and t.startswith(q):
            return 1
        if q and q in t:
            return 2
        return 3

    merged.sort(key=lambda r: (rank(r), 0 if r["pages"] else 1))
    return merged


# ============================================================================
# EXCEL İÇE / DIŞA AKTARMA (openpyxl gerekir)
# ============================================================================

def _require_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        messagebox.showerror(APP_NAME, T("need_openpyxl"))
        return False


_SCALE_SUFFIX = re.compile(r"\s*\(\s*\d+\s*-\s*(\d+)\s*\)\s*$")


def _status_lookup():
    """Durum adı (6 dilden biri ya da kod) → kod."""
    lk = {code: code for code in STATUSES}
    for d in I18N.values():
        for code in STATUSES:
            lk.setdefault(d["st_" + code].casefold(), code)
    return lk


def import_excel(path):
    """Beklenen düzen: Kitap | Yazar | Tür | Sayfa | ölçüt sütunları… | Skor
    (başlıklar 6 dilden herhangi birinde olabilir). Ağırlıklar ikinci sayfadan
    (A: ölçüt adı, B: ağırlık) okunur."""
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header = [str(c.value or "").strip() for c in ws[1]]
    low = [h.lower() for h in header]

    def find(aliases):
        return next((i for i, h in enumerate(low) if h in aliases), None)

    ti, ai, gi, pi = find(_AL_TITLE), find(_AL_AUTHOR), find(_AL_GENRE), find(_AL_PAGES)
    si = next((i for i, h in enumerate(low) if any(h.startswith(a) for a in _AL_SCORE)), None)
    sti = find(_AL_STATUS)
    if ti is None:
        raise ValueError("Title column not found / 'Kitap Adı'")

    known = {i for i in (ti, ai, gi, pi, si, sti) if i is not None}
    crit_cols = [i for i, h in enumerate(header) if h and i not in known]
    if not crit_cols:
        raise ValueError("No scoring-criteria columns found")

    weights = {}
    if len(wb.worksheets) > 1:
        for row in wb.worksheets[1].iter_rows(min_row=2, max_col=2, values_only=True):
            if (row[0] and isinstance(row[1], (int, float)) and not isinstance(row[1], bool)
                    and math.isfinite(row[1]) and row[1] >= 0):
                weights[str(row[0]).strip().lower()] = float(row[1])

    scale_max = DEFAULT_SCALE_MAX
    criteria = []
    for i in crit_cols:
        raw = header[i]
        m = _SCALE_SUFFIX.search(raw)
        if m:
            scale_max = int(m.group(1))
        name = _SCALE_SUFFIX.sub("", raw).strip()
        criteria.append({"id": new_id(), "name": name, "col": i,
                         "weight": weights.get(name.lower(), 1.0)})

    books = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        title = row[ti] if ti < len(row) else None
        if not title or not str(title).strip():
            # başlığı boş ama başka sütunu dolu satır (tablonun boş satırından eklenmiş olabilir) kaybolmasın
            others = [i for i in (ai, gi, pi, *crit_cols) if i is not None and i < len(row)]
            if not any(row[i] not in (None, "") for i in others):
                continue
            title = ""

        def cell(idx):
            return row[idx] if idx is not None and idx < len(row) else None

        pages = cell(pi)
        try:
            pages = int(float(pages)) if pages not in (None, "") else None
        except (TypeError, ValueError, OverflowError):
            pages = None
        b = new_book(str(title).strip(), str(cell(ai) or "").strip(), str(cell(gi) or "").strip(), pages)
        stv = str(cell(sti) or "").strip().casefold()
        if stv:
            code = _status_lookup().get(stv)
            if code and code != "todo":
                b["status"] = code
        for c in criteria:
            v = cell(c["col"])
            if isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v):
                v = min(v, scale_max)   # arayüzdeki ayar penceresiyle aynı kural: üst sınırı aşan puan kırpılır
                b["scores"][c["id"]] = float(v) if v != int(v) else int(v)
        books.append(b)

    for c in criteria:
        c.pop("col")
    return {"scale_max": scale_max, "criteria": tag_builtin_criteria(criteria), "books": books}


def _xl_text(cell, value):
    """Metin hücresi yaz; '=' ile başlayan değer (ör. kitap adı) formül olarak yorumlanmasın."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def export_excel(data, path):
    """Formüllü, renk skalalı ve veri çubuklu çıktı."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    cs, books, smax = data["criteria"], data["books"], data["scale_max"]
    n = len(cs)
    score_col = 5 + n
    wb = Workbook()
    ws = wb.active
    sheet2 = T("xl_sheet_settings")[:31]
    main_name = re.sub(r"[\\/*?:\[\]]+", " ", str(data.get("name") or "")).strip()[:31] or T("xl_sheet_main")[:31]
    if main_name.casefold() == sheet2.casefold():
        main_name = main_name[:28] + " 1"
    ws.title = main_name
    st = wb.create_sheet(sheet2)
    ref = "'" + sheet2.replace("'", "''") + "'"

    head_fill = PatternFill("solid", fgColor=THEMES["cream"]["header_bg"].lstrip("#"))
    head_font = Font(bold=True, color="FFFFFF")
    headers = [T("col_title"), T("col_author"), T("col_genre"), T("col_pages")] + \
              [f"{c['name']} (1-{smax})" for c in cs] + [T("col_score"), T("col_status")]
    for i, h in enumerate(headers, start=1):
        cell = _xl_text(ws.cell(row=1, column=i), h)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 45

    st.cell(row=1, column=1, value=T("xl_weights_head")).fill = head_fill
    st.cell(row=1, column=2, value=T("xl_weight")).fill = head_fill
    st["A1"].font = st["B1"].font = head_font
    for i, c in enumerate(cs):
        _xl_text(st.cell(row=2 + i, column=1), c["name"])
        wc = st.cell(row=2 + i, column=2, value=c["weight"])
        wc.fill = PatternFill("solid", fgColor="FFF6DC")
        wc.font = Font(bold=True)
    st.cell(row=2 + n, column=1, value=T("xl_total")).font = Font(bold=True)
    st.cell(row=2 + n, column=2, value=f"=SUM(B2:B{1 + n})").font = Font(bold=True)
    st.cell(row=2 + n, column=3, value=T("xl_note"))
    st.column_dimensions["A"].width = 36
    st.column_dimensions["C"].width = 60

    for r, b in enumerate(books, start=2):
        _xl_text(ws.cell(row=r, column=1), b["title"]).font = Font(bold=True)
        _xl_text(ws.cell(row=r, column=2), b["author"])
        _xl_text(ws.cell(row=r, column=3), b["genre"])
        ws.cell(row=r, column=4, value=b["pages"])
        for i, c in enumerate(cs):
            v = b["scores"].get(c["id"])
            if v not in (None, ""):
                ws.cell(row=r, column=5 + i, value=v)
        num = "+".join(f'IF({L(5+i)}{r}="",0,{L(5+i)}{r}*{ref}!$B${2+i})' for i in range(n))
        den = "+".join(f'IF({L(5+i)}{r}="",0,{ref}!$B${2+i})' for i in range(n))
        f = ws.cell(row=r, column=score_col, value=f"=IF(({den})=0,0,ROUND(({num})/({den}),2))")
        f.font = Font(bold=True)
        _xl_text(ws.cell(row=r, column=score_col + 1), T("st_" + book_status(b)))

    last = len(books) + 101   # sonradan Excel'de eklenecek satırlar için pay
    sc = L(score_col)
    ws.conditional_formatting.add(
        f"E2:{L(4+n)}{last}",
        ColorScaleRule(start_type="num", start_value=1, start_color="FBE5E1",
                       mid_type="num", mid_value=(1 + smax) / 2, mid_color="FFF6DC",
                       end_type="num", end_value=smax, end_color="CDEBD6"))
    ws.conditional_formatting.add(
        f"{sc}2:{sc}{last}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=smax, color="7FA6D9"))

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 17
    ws.column_dimensions["D"].width = 14
    for i in range(n):
        ws.column_dimensions[L(5 + i)].width = 15
    ws.column_dimensions[sc].width = 14
    ws.column_dimensions[L(score_col + 1)].width = 16
    ws.freeze_panes = "B2"
    wb.save(path)


# ============================================================================
# YARDIMCI: pencere ortalama
# ============================================================================

def center_on_screen(win):
    try:
        win.update_idletasks()
        w, h = win.winfo_reqwidth(), win.winfo_reqheight()
        x = max(0, (win.winfo_screenwidth() - w) // 2)
        y = max(0, (win.winfo_screenheight() - h) // 3)
        win.geometry(f"+{x}+{y}")
    except tk.TclError:
        pass


def center_on_parent(win, parent, w, h):
    try:
        parent.update_idletasks()
        x = parent.winfo_rootx() + max(0, (parent.winfo_width() - w) // 2)
        y = parent.winfo_rooty() + max(0, (parent.winfo_height() - h) // 3)
    except tk.TclError:
        x = y = 100
    win.geometry(f"{w}x{h}+{x}+{y}")


# ============================================================================
# İLK AÇILIŞ: DİL SEÇİMİ
# ============================================================================

class LanguageDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app.root)
        t = app.theme
        self.result = None
        self.title(APP_NAME)
        self.configure(bg=t["bg"], padx=26, pady=22)
        self.resizable(False, False)
        try:
            self.iconbitmap(resource_path("next_book.ico"))
        except Exception:
            pass
        fam = app.default_family()
        tk.Label(self, text="Next Book", font=("Georgia", 24, "bold"),
                 bg=t["bg"], fg=t["title_fg"]).pack()
        tk.Label(self, text="Dil seçin · Choose your language · Выберите язык\n"
                            "Sprache wählen · Choisir la langue · 选择语言",
                 font=(fam, 10), bg=t["bg"], fg=t["muted"], justify="center").pack(pady=(4, 16))
        grid = tk.Frame(self, bg=t["bg"])
        grid.pack()
        for i, (code, name) in enumerate(LANGS):
            img = app.flag_image(code, 60, 40)
            btn = tk.Button(grid, image=img, text=name, compound="top", font=(fam, 11),
                            bd=0, relief="flat", bg=t["button"], activebackground=t["button_active"],
                            fg=t["text"], activeforeground=t["text"], padx=16, pady=10,
                            cursor="hand2", command=lambda c=code: self._pick(c))
            btn.grid(row=i // 3, column=i % 3, padx=6, pady=6)
        self.protocol("WM_DELETE_WINDOW", lambda: self._pick("en"))
        center_on_screen(self)
        try:
            self.grab_set()
        except tk.TclError:
            pass
        self.focus_force()

    def _pick(self, code):
        self.result = code
        self.destroy()


# ============================================================================
# HAKKINDA KARTI
# ============================================================================

class ManualBookDialog(tk.Toplevel):
    """Elle ekleme: kitap adı, yazar, tür, sayfa sayısı. En az biri dolu olmalı (yalnızca yazar da olabilir)."""

    KEYS = ("title", "author", "genre", "pages")

    def __init__(self, app):
        super().__init__(app.root)
        t = app.theme
        self.app, self.result = app, None
        self.title(T("btn_manual").lstrip("+ ").strip())
        self.configure(bg=t["bg"], padx=22, pady=18)
        self.transient(app.root)
        self.resizable(False, False)
        try:
            self.iconbitmap(resource_path("next_book.ico"))
        except Exception:
            pass
        body = ttk.Frame(self)
        body.pack(fill="x")
        body.columnconfigure(0, weight=1)
        self.entries = {}
        for r, key in enumerate(self.KEYS):
            ttk.Label(body, text=T("col_" + key)).grid(row=r * 2, column=0, sticky="w", pady=(0 if r == 0 else 10, 2))
            e = ttk.Entry(body, width=14 if key == "pages" else 46)
            e.grid(row=r * 2 + 1, column=0, sticky="w" if key == "pages" else "ew")
            e.bind("<Return>", lambda _e, k=key: self._enter(k))
            if key in ("author", "genre"):
                attach_autocomplete(e, lambda k=key: suggest_values(app.all_books(), k))
            self.entries[key] = e
        self.msg = ttk.Label(self, text="", foreground="#FF8A80" if t["dark"] else "#B3261E")
        self.msg.pack(fill="x", pady=(8, 0))
        btns = ttk.Frame(self)
        btns.pack(fill="x", pady=(10, 0))
        ttk.Button(btns, text=T("btn_add_to_list"), style="Accent.TButton", command=self.submit).pack(side="right")
        ttk.Button(btns, text=T("btn_cancel"), command=self.destroy).pack(side="right", padx=(0, 8))
        self.bind("<Escape>", lambda _e: self.destroy())
        try:
            self.update_idletasks()
            center_on_parent(self, app.root, self.winfo_reqwidth(), self.winfo_reqheight())
        except tk.TclError:
            pass
        try:
            self.grab_set()
        except tk.TclError:
            pass
        self.entries["title"].focus_set()

    def _enter(self, key):
        e = self.entries[key]
        if e.selection_present():      # otomatik tamamlama önerisi varsa önce onu kabul et
            e.icursor("end")
            e.selection_clear()
            return "break"
        self.submit()
        return "break"

    def submit(self):
        title = self.entries["title"].get().strip()
        if not any(self.entries[k].get().strip() for k in self.KEYS):
            self.bell()
            self.entries["title"].focus_set()
            return
        raw = self.entries["pages"].get().strip()
        pages = None
        if raw:
            try:
                pages = int(float(raw.replace(",", ".")))
                if pages < 0:
                    raise ValueError
            except (ValueError, OverflowError):
                self.bell()
                self.msg.config(text=T("bad_pages"))
                self.entries["pages"].focus_set()
                return
        self.result = {"title": title, "author": self.entries["author"].get().strip(),
                       "genre": self.entries["genre"].get().strip(), "pages": pages}
        self.destroy()


class AboutDialog(tk.Toplevel):
    """Hakkında kartı: tema renklerine uyar, 6 dilde."""

    WIDTH = 440

    def __init__(self, app):
        super().__init__(app.root)
        t = app.theme
        fam, size = app.default_family(), app.font_size
        small = max(8, size - 1)
        self.title(T("about_title"))
        self.configure(bg=t["bg"], padx=28, pady=24)
        self.transient(app.root)
        self.resizable(False, False)
        try:
            self.iconbitmap(resource_path("next_book.ico"))
        except Exception:
            pass

        tk.Label(self, text=APP_NAME.upper(), font=(fam, small, "bold"),
                 bg=t["bg"], fg=t["muted"], anchor="w").pack(fill="x")
        tk.Label(self, text=T("about_title"), font=(fam, size + 10, "bold"),
                 bg=t["bg"], fg=t["title_fg"], anchor="w").pack(fill="x", pady=(2, 6))
        tk.Label(self, text=f"{APP_NAME} v{APP_VERSION}", font=(fam, small, "bold"),
                 bg=t["bg"], fg=t["accent"], anchor="w").pack(fill="x", pady=(0, 14))
        for key in ("about_hello", "about_privacy"):
            tk.Label(self, text=T(key), font=app.font_normal, bg=t["bg"], fg=t["text"],
                     justify="left", anchor="w", wraplength=self.WIDTH).pack(fill="x", pady=(0, 10))

        tk.Label(self, text=f"{T('about_data')}: {DATA_DIR}", font=(fam, small), bg=t["bg"], fg=t["muted"],
                 justify="left", anchor="w", wraplength=self.WIDTH).pack(fill="x", pady=(4, 0))

        btns = tk.Frame(self, bg=t["bg"])
        btns.pack(fill="x", pady=(18, 0))
        ttk.Button(btns, text=T("btn_close"), style="Accent.TButton", command=self.destroy).pack(side="right")
        self.bind("<Escape>", lambda _e: self.destroy())
        try:
            self.update_idletasks()
            center_on_parent(self, app.root, self.winfo_reqwidth(), self.winfo_reqheight())
        except tk.TclError:
            pass
        try:
            self.grab_set()
        except tk.TclError:
            pass
        self.focus_set()


# ============================================================================
# DUYURU PANOSU (başlıktaki duyuru kutusuna tıklayınca açılır)
# ============================================================================

class AnnouncementBoard(tk.Toplevel):
    def __init__(self, app, items):
        super().__init__(app.root)
        self.app = app
        t = app.theme
        fam, size = app.default_family(), app.font_size
        small = (fam, max(8, size - 1), "bold")
        self.labels = []
        self.title(T("ann_board_title"))
        self.configure(bg=t["bg"])
        self.transient(app.root)
        self.minsize(420, 320)
        try:
            self.iconbitmap(resource_path("next_book.ico"))
        except Exception:
            pass

        head = tk.Frame(self, bg=t["bg"])
        head.pack(side="top", fill="x", padx=20, pady=(16, 8))
        tk.Label(head, text=T("ann_board_title"), font=(fam, size + 7, "bold"),
                 bg=t["bg"], fg=t["title_fg"]).pack(side="left")
        if items:
            tk.Label(head, text=str(len(items)), font=small, bg=t["accent"], fg=t["accent_fg"],
                     padx=8, pady=1).pack(side="left", padx=10)

        foot = tk.Frame(self, bg=t["bg"])
        foot.pack(side="bottom", fill="x", padx=20, pady=(8, 16))
        ttk.Button(foot, text=T("btn_close"), style="Accent.TButton", command=self.destroy).pack(side="right")

        mid = tk.Frame(self, bg=t["bg"])
        mid.pack(side="top", fill="both", expand=True, padx=(20, 8))
        self.canvas = tk.Canvas(mid, bg=t["bg"], highlightthickness=0, bd=0)
        sb = ttk.Scrollbar(mid, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=sb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.inner = tk.Frame(self.canvas, bg=t["bg"])
        self.win_id = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda _e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", self._on_canvas_resize)
        self.bind("<MouseWheel>", self._on_wheel)
        self.bind("<Escape>", lambda _e: self.destroy())

        if not items:
            msg = tk.Label(self.inner, text=T("ann_empty"), font=app.font_normal, bg=t["bg"], fg=t["muted"],
                           justify="left", anchor="w", wraplength=480)
            msg.pack(fill="x", pady=20)
            self.labels.append(msg)
        icon_px = max(14, round(app.font_normal.metrics("linespace") * 0.95))
        for a in items:
            kind = a["type"] if a["type"] in ANN_TYPES else "info"
            fill, glyph = ann_icon_colors(t, kind)
            card = tk.Frame(self.inner, bg=t["field"], highlightthickness=1, highlightbackground=t["border"])
            card.pack(fill="x", pady=(0, 10), padx=(0, 6))
            tk.Frame(card, bg=fill, width=5).pack(side="left", fill="y")
            body = tk.Frame(card, bg=t["field"])
            body.pack(side="left", fill="both", expand=True, padx=14, pady=10)
            row = tk.Frame(body, bg=t["field"])
            row.pack(fill="x")
            tk.Label(row, image=app.ann_icon_img(kind, fill, glyph, t["field"], icon_px),
                     bg=t["field"], bd=0).pack(side="left")
            tk.Label(row, text=T("ann_type_" + kind), font=small, bg=t["field"], fg=t["muted"],
                     anchor="w").pack(side="left", padx=(8, 0))
            msg = tk.Label(body, text=a["text"], font=app.font_normal, bg=t["field"], fg=t["field_fg"],
                           justify="left", anchor="w", wraplength=480)
            msg.pack(fill="x", pady=(4, 0))
            self.labels.append(msg)
            if a.get("url"):
                ttk.Button(body, text="↗  " + T("ann_open"),
                           command=lambda u=a["url"], tx=a["text"]: app.open_url_confirmed(tx, u, parent=self)
                           ).pack(anchor="w", pady=(8, 0))

        center_on_parent(self, app.root, 600, 520)
        try:
            self.grab_set()
        except tk.TclError:
            pass
        self.focus_set()

    def _on_canvas_resize(self, e):
        self.canvas.itemconfigure(self.win_id, width=e.width)
        for lbl in self.labels:
            lbl.configure(wraplength=max(200, e.width - 70))

    def _on_wheel(self, e):
        self.canvas.yview_scroll(-1 if e.delta > 0 else 1, "units")


# ============================================================================
# YAZI TİPİ PENCERESİ
# ============================================================================

class FontDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app.root)
        self.app = app
        t = app.theme
        self.title(T("font_title"))
        self.configure(bg=t["bg"])
        self.transient(app.root)
        self.resizable(False, False)
        installed = {f for f in tkfont.families() if f and not f.startswith("@")}
        families = [f for f in FONT_CHOICES if f in installed]
        dflt = app.language_default_family()
        if dflt in installed and dflt not in families:   # ör. Çince arayüzün varsayılanı
            families.insert(0, dflt)
        if not families:   # Windows dışında hiçbiri yoksa tüm listeyi göster
            families = sorted(installed, key=str.casefold)
        self.fam_var = tk.StringVar(value=app.default_family())
        self.size_var = tk.StringVar(value=str(app.font_size))

        body = ttk.Frame(self, padding=16)
        body.pack(fill="both", expand=True)
        ttk.Label(body, text=T("font_family")).grid(row=0, column=0, sticky="w")
        self.combo = ttk.Combobox(body, values=families, textvariable=self.fam_var, state="readonly", width=34)
        self.combo.grid(row=1, column=0, sticky="ew", pady=(2, 10))
        ttk.Label(body, text=T("font_size")).grid(row=0, column=1, sticky="w", padx=(12, 0))
        self.spin = ttk.Spinbox(body, from_=8, to=18, width=5, textvariable=self.size_var, command=self._preview)
        self.spin.grid(row=1, column=1, padx=(12, 0), pady=(2, 10))

        self.preview_font = tkfont.Font(family=app.default_family(), size=app.font_size)
        self.preview = tk.Label(body, text=T("font_preview"), font=self.preview_font, bg=t["field"],
                                fg=t["field_fg"], relief="solid", bd=1, padx=12, pady=16, width=44)
        self.preview.grid(row=2, column=0, columnspan=2, sticky="ew")

        btns = ttk.Frame(body)
        btns.grid(row=3, column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(btns, text=T("font_reset"), command=self._reset).pack(side="left")
        ttk.Button(btns, text=T("btn_cancel"), command=self.destroy).pack(side="left", padx=(14, 6))
        ttk.Button(btns, text=T("btn_apply"), style="Accent.TButton", command=self._apply).pack(side="left")

        self.combo.bind("<<ComboboxSelected>>", lambda _e: self._preview())
        self.spin.bind("<KeyRelease>", lambda _e: self._preview())
        center_on_screen(self)
        try:
            self.grab_set()
        except tk.TclError:
            pass

    def _size(self):
        try:
            return max(8, min(18, int(float(self.size_var.get()))))
        except ValueError:
            return self.app.font_size

    def _preview(self):
        self.preview_font.configure(family=self.fam_var.get(), size=self._size())

    def _reset(self):
        self.fam_var.set(self.app.language_default_family())
        self.size_var.set("10")
        self._preview()

    def _apply(self):
        fam = self.fam_var.get()
        default = self.app.language_default_family()
        self.app.set_font(None if fam == default else fam, self._size())
        self.destroy()


# ============================================================================
# PUANLAMA AYARLARI PENCERESİ (kaydırılabilir — istediğin kadar ölçüt)
# ============================================================================

class SettingsDialog(tk.Toplevel):
    def __init__(self, app):
        super().__init__(app.root)
        self.app = app
        t = app.theme
        self.title(T("st_title"))
        self.configure(bg=t["bg"])
        self.transient(app.root)
        self.minsize(520, 380)
        self.rows = []

        # ---- üst: açıklama
        head = ttk.Frame(self, padding=(16, 14, 16, 6))
        head.pack(side="top", fill="x")
        ttk.Label(head, text=T("st_heading"), font=app.font_head).pack(anchor="w")
        self.desc = ttk.Label(head, text=T("st_desc"), foreground=t["muted"], justify="left", wraplength=500)
        self.desc.pack(anchor="w", pady=(4, 0))

        # ---- alt: toplam + kaydet/iptal (gövdeden ÖNCE paketlenir ki hep görünsün)
        foot = ttk.Frame(self, padding=(16, 8, 16, 14))
        foot.pack(side="bottom", fill="x")
        self.total_lbl = ttk.Label(foot, foreground=t["muted"])     # satırlardan ÖNCE oluşur
        self.total_lbl.pack(side="left")
        ttk.Button(foot, text=T("btn_cancel"), command=self.destroy).pack(side="right", padx=(8, 0))
        ttk.Button(foot, text=T("btn_save"), style="Accent.TButton", command=self.save).pack(side="right")

        ctl = ttk.Frame(self, padding=(16, 6, 16, 0))
        ctl.pack(side="bottom", fill="x")
        ttk.Button(ctl, text=T("st_add"), command=self._add_new).pack(side="left")
        ttk.Label(ctl, text="    " + T("st_scale")).pack(side="left")
        self.scale_var = tk.StringVar(value=str(app.data["scale_max"]))
        ttk.Spinbox(ctl, from_=2, to=10, width=4, textvariable=self.scale_var).pack(side="left", padx=4)

        # ---- orta: kaydırılabilir ölçüt listesi
        mid = ttk.Frame(self, padding=(16, 4, 10, 0))
        mid.pack(side="top", fill="both", expand=True)
        self.canvas = tk.Canvas(mid, bg=t["bg"], highlightthickness=0, bd=0)
        sb = ttk.Scrollbar(mid, orient="vertical", command=self.canvas.yview)
        self.canvas.configure(yscrollcommand=sb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.inner = ttk.Frame(self.canvas)
        self.win_id = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda _e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", self._on_canvas_resize)
        self.bind("<MouseWheel>", self._on_wheel)

        self.inner.columnconfigure(0, weight=1)
        ttk.Label(self.inner, text=T("st_name"), foreground=t["muted"]).grid(row=0, column=0, sticky="w")
        ttk.Label(self.inner, text=T("st_weight"), foreground=t["muted"]).grid(row=0, column=1, sticky="w", padx=6)

        for c in app.data["criteria"]:
            self._add_row(c["id"], c["name"], c["weight"], builtin=c.get("builtin"))

        center_on_parent(self, app.root, 580, 540)
        try:
            self.grab_set()
        except tk.TclError:
            pass
        self.focus_set()

    # ---- kaydırma
    def _on_canvas_resize(self, e):
        self.canvas.itemconfigure(self.win_id, width=e.width)
        self.desc.configure(wraplength=max(300, e.width - 20))

    def _on_wheel(self, e):
        self.canvas.yview_scroll(-1 if e.delta > 0 else 1, "units")

    # ---- satırlar
    def _add_new(self):
        self._add_row(new_id(), "", 1.0, focus=True)
        self.after(60, lambda: self.canvas.yview_moveto(1.0))

    def _add_row(self, cid, name, weight, focus=False, builtin=None):
        name_var = tk.StringVar(value=name)
        weight_var = tk.StringVar(value=fmt_num(weight))
        weight_var.trace_add("write", lambda *_: self._update_total())
        e1 = ttk.Entry(self.inner, textvariable=name_var)
        e2 = ttk.Entry(self.inner, textvariable=weight_var, width=7)
        pct = ttk.Label(self.inner, width=5, anchor="e", foreground=self.app.theme["muted"])
        row = {"id": cid, "name": name_var, "weight": weight_var, "pct": pct, "builtin": builtin}
        btn = ttk.Button(self.inner, text="✕", width=3, command=lambda: self._remove_row(row))
        row["widgets"] = [e1, e2, pct, btn]
        self.rows.append(row)
        self._regrid()
        self._update_total()
        if focus:
            e1.focus_set()

    def _remove_row(self, row):
        if len(self.rows) <= 1:
            messagebox.showinfo(APP_NAME, T("st_need_one"), parent=self)
            return
        for w in row["widgets"]:
            w.destroy()
        self.rows.remove(row)
        self._regrid()
        self._update_total()

    def _regrid(self):
        for i, row in enumerate(self.rows, start=1):
            e1, e2, pct, btn = row["widgets"]
            e1.grid(row=i, column=0, sticky="ew", pady=3)
            e2.grid(row=i, column=1, padx=6, pady=3)
            pct.grid(row=i, column=2, pady=3)
            btn.grid(row=i, column=3, padx=(4, 0), pady=3)

    def _update_total(self):
        vals = []
        for row in self.rows:
            try:
                v = parse_float(row["weight"].get())
                vals.append(max(0.0, v) if math.isfinite(v) else 0.0)
            except ValueError:
                vals.append(0.0)
        total = sum(vals)
        self.total_lbl.config(text=T("st_total", total=f"{total:.2f}"))
        for row, v in zip(self.rows, vals):
            row["pct"].config(text=f"{v / total * 100:.0f}%" if total > 0 else "")

    # ---- kaydet
    def save(self):
        criteria = []
        for row in self.rows:
            name = row["name"].get().strip()
            if not name:
                messagebox.showwarning(APP_NAME, T("st_name_required"), parent=self)
                return
            try:
                w = parse_float(row["weight"].get())
                if not math.isfinite(w) or w < 0:
                    raise ValueError
            except ValueError:
                messagebox.showwarning(APP_NAME, T("st_bad_weight", name=name), parent=self)
                return
            c = {"id": row["id"], "name": name, "weight": w}
            if row.get("builtin") and name == T(row["builtin"]):   # adı değiştirilmediyse dil değişince çevrilsin
                c["builtin"] = row["builtin"]
            criteria.append(c)
        if len({c["name"].lower() for c in criteria}) != len(criteria):
            messagebox.showwarning(APP_NAME, T("st_dup_name"), parent=self)
            return
        try:
            smax = int(self.scale_var.get())
            if not 2 <= smax <= 10:
                raise ValueError
        except ValueError:
            messagebox.showwarning(APP_NAME, T("st_bad_scale"), parent=self)
            return
        self.app.apply_settings(criteria, smax)
        self.destroy()


# ============================================================================
# ANA UYGULAMA
# ============================================================================

class NextBookApp:
    FIXED = ["title", "author", "genre", "pages"]
    ICON_PAD = 3   # #0 hücresinde rozetin solunda kalan sabit boşluk (px)
    SWATCH = 26

    def __init__(self, root):
        self.root = root
        self.settings = load_settings()
        self.settings.pop("dismissed_ann", None)   # duyurular artık kapatılamıyor; eski kayıt gereksiz
        self.theme_name = self.settings.get("theme") if self.settings.get("theme") in THEMES else DEFAULT_THEME
        self.theme = THEMES[self.theme_name]
        fam = self.settings.get("font_family")
        self.font_family = fam if fam in FONT_CHOICES else None   # listede olmayan eski seçim varsayılana döner
        try:
            self.font_size = max(8, min(18, int(self.settings.get("font_size", 10))))
        except (TypeError, ValueError):
            self.font_size = 10
        set_language(self.settings.get("lang") if self.settings.get("lang") in I18N else "en")

        self.grid_on = bool(self.settings.get("grid", True))   # ızgara çizgileri açık mı
        self._grid_job = None
        self._hlines, self._vlines = [], []
        self._created = None
        self.flag_cache = {}
        self.icon_cache = {}
        self.results = []
        self.q = queue.Queue()
        self.ann_feed = None
        self.ann_idx = 0
        self.ann_cur = None
        self._ann_job = None
        self._ann_over = False
        self.edit = None
        self.sort_col = None
        self.sort_rev = False
        self.lang_popup = None
        self._hover_cid = None
        self.title_var = tk.StringVar()
        self.author_var = tk.StringVar()

        root.title(f"{APP_NAME} {APP_VERSION}")
        try:
            root.iconbitmap(resource_path("next_book.ico"))
        except Exception:
            pass   # simge yoksa ya da platform .ico desteklemiyorsa sorun değil
        root.geometry("1220x760")      # pencere küçültülürse dönülecek boyut
        root.minsize(900, 560)
        try:                            # tam ekran boyutunda aç
            root.state("zoomed")
        except tk.TclError:
            try:
                root.attributes("-zoomed", True)
            except tk.TclError:
                pass

        self._make_fonts()
        self._setup_style()
        if self.settings.get("lang") not in I18N:
            self._first_run_language()
        self.store = load_data(current_language())
        self.data = active_list(self.store)   # self.data her zaman AÇIK sekmenin listesi
        self._drag = None
        self.tab_items = []
        if self._localize_all():   # varsayılan ölçütler / ilk liste adı arayüz diline uysun
            self.save()

        self._build_ui()
        self.rebuild_columns()
        self.refresh_table()
        self._poll_queue()
        self._announce_tick()
        if LOAD_NOTICE:
            path = LOAD_NOTICE[0]
            root.after(400, lambda: messagebox.showwarning(APP_NAME, T("data_recovered", path=path)))
        root.bind("<Button-1>", self._on_root_click, add="+")
        root.protocol("WM_DELETE_WINDOW", self.on_close)

    def _localize_all(self):
        changed = False
        for l in self.store["lists"]:
            if localize_criteria(l["criteria"]):
                changed = True
            if l.get("name_builtin") and l["name"] != T("tab_default"):
                l["name"] = T("tab_default")
                changed = True
        return changed

    def all_books(self):
        return [b for l in self.store["lists"] for b in l["books"]]

    # ---------------------------------------------------------------- ayarlar
    def _persist_settings(self):
        self.settings.update({"lang": current_language(), "theme": self.theme_name,
                              "font_family": self.font_family, "font_size": self.font_size,
                              "grid": self.grid_on})
        save_settings(self.settings)

    def _first_run_language(self):
        self.root.withdraw()
        dlg = LanguageDialog(self)
        self.root.wait_window(dlg)
        set_language(dlg.result or "en")
        self._apply_fonts()
        self._persist_settings()
        self.root.deiconify()

    def language_default_family(self):
        return "Microsoft YaHei UI" if current_language() == "zh" else "Segoe UI"

    def default_family(self):
        return self.font_family or self.language_default_family()

    def flag_image(self, code, w, h):
        key = (code, w, h)
        if key not in self.flag_cache:
            self.flag_cache[key] = make_flag_image(code, w, h)
        return self.flag_cache[key]

    # ---------------------------------------------------------------- yazı tipleri
    def _make_fonts(self):
        fam, size = self.default_family(), self.font_size
        self.font_normal = tkfont.Font(family=fam, size=size)
        self.font_head = tkfont.Font(family=fam, size=size, weight="bold")

    def _apply_fonts(self):
        fam, size = self.default_family(), self.font_size
        for f in (self.font_normal, self.font_head):
            f.configure(family=fam, size=size)
        for name in ("TkDefaultFont", "TkTextFont", "TkMenuFont", "TkHeadingFont"):
            try:
                tkfont.nametofont(name).configure(family=fam, size=size)
            except tk.TclError:
                pass
        style = ttk.Style()
        style.configure("Treeview", font=self.font_normal, rowheight=max(24, round(size * 2.8)))
        style.configure("Treeview.Heading", font=self.font_head)
        style.configure("TLabelframe.Label", font=self.font_head)
        style.configure("Accent.TButton", font=self.font_head)

    def set_font(self, family, size):
        self.font_family = family
        self.font_size = size
        self._persist_settings()
        self._apply_fonts()
        self.rebuild_ui()

    # ---------------------------------------------------------------- tema / stil
    def _setup_style(self):
        t = self.theme
        self.root.configure(bg=t["bg"])
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(".", background=t["bg"], foreground=t["text"])
        style.configure("TFrame", background=t["bg"])
        style.configure("TLabel", background=t["bg"], foreground=t["text"])
        style.configure("TLabelframe", background=t["bg"], bordercolor=t["border"])
        style.configure("TLabelframe.Label", background=t["bg"], foreground=t["label_frame"])
        style.configure("TButton", background=t["button"], foreground=t["text"],
                        bordercolor=t["border"], padding=(10, 4))
        style.map("TButton", background=[("active", t["button_active"]), ("disabled", t["button_disabled"])])
        style.configure("Tool.TButton", padding=(8, 3))
        style.configure("Grid.Toolbutton", background=t["button"], foreground=t["text"],
                        bordercolor=t["border"], padding=(10, 4))
        style.map("Grid.Toolbutton", background=[("selected", t["accent"]), ("active", t["button_active"])],
                  foreground=[("selected", t["accent_fg"])])
        style.configure("Accent.TButton", background=t["accent"], foreground=t["accent_fg"], padding=(16, 4))
        style.map("Accent.TButton", background=[("active", t["accent_active"]), ("disabled", t["button_disabled"])])
        style.configure("TEntry", fieldbackground=t["field"], foreground=t["field_fg"],
                        insertcolor=t["field_fg"], bordercolor=t["border"], padding=3)
        style.configure("TSpinbox", fieldbackground=t["field"], foreground=t["field_fg"],
                        arrowcolor=t["text"], bordercolor=t["border"], background=t["button"])
        style.configure("TCombobox", fieldbackground=t["field"], foreground=t["field_fg"],
                        arrowcolor=t["text"], bordercolor=t["border"], background=t["button"])
        style.map("TCombobox", fieldbackground=[("readonly", t["field"])],
                  foreground=[("readonly", t["field_fg"])],
                  selectbackground=[("readonly", t["field"])], selectforeground=[("readonly", t["field_fg"])])
        for orient in ("Vertical", "Horizontal"):
            style.configure(f"{orient}.TScrollbar", background=t["button"], troughcolor=t["bg"],
                            bordercolor=t["border"], arrowcolor=t["text"])
        style.configure("Treeview", background=t["tree_bg"], fieldbackground=t["tree_bg"],
                        foreground=t["text"], borderwidth=0)
        style.configure("Treeview.Heading", background=t["header_bg"], foreground=t["header_fg"],
                        relief="flat", padding=6)
        style.map("Treeview.Heading", background=[("active", t["header_active"])])
        try:   # satır başındaki "ağaç oku" boşluğunu kaldır: rozet o boşluk yüzünden sağa kayıyordu
            # Treeitem.focus de çıkarıldı: boş (rozetsiz) #0 hücresinde noktalı odak dikdörtgeni çiziyordu
            style.layout("Treeview.Item", [("Treeitem.padding", {"sticky": "nswe", "children": [
                ("Treeitem.image", {"side": "left", "sticky": ""}),
                ("Treeitem.text", {"side": "left", "sticky": ""})]})])
            style.configure("Treeview.Item", padding=(self.ICON_PAD, 0, 0, 0))
        except tk.TclError:
            pass
        style.map("Treeview", background=[("selected", t["select_bg"])],
                  foreground=[("selected", t["select_fg"])])
        try:   # Combobox açılır listesi
            self.root.option_add("*TCombobox*Listbox.background", t["field"])
            self.root.option_add("*TCombobox*Listbox.foreground", t["field_fg"])
            self.root.option_add("*TCombobox*Listbox.selectBackground", t["select_bg"])
            self.root.option_add("*TCombobox*Listbox.selectForeground", t["select_fg"])
        except tk.TclError:
            pass
        self._apply_fonts()

    def set_theme(self, name):
        if name not in THEMES or name == self.theme_name:
            return
        self.theme_name = name
        self.theme = THEMES[name]
        self._persist_settings()
        self._setup_style()
        self.rebuild_ui()

    def set_lang(self, code):
        if code not in I18N:
            return
        set_language(code)
        if self._localize_all():
            self.save()
        self._persist_settings()
        self._apply_fonts()
        self.rebuild_ui()

    def rebuild_ui(self):
        """Tema / dil / yazı tipi değişince arayüzü baştan kur (veri aynı kalır)."""
        self.finish_edit(False)
        self._cancel_ann_job()
        if self.lang_popup is not None:
            try:
                self.lang_popup.destroy()
            except tk.TclError:
                pass
            self.lang_popup = None
        for w in list(self.root.winfo_children()):
            w.destroy()
        self._hover_cid = None
        self._grid_job = None
        self._build_ui()
        self.rebuild_columns()
        self.refresh_table()
        if self.results:
            self.show_results(self.results)

    # ---------------------------------------------------------------- arayüz
    def _build_ui(self):
        r, t = self.root, self.theme
        self._hlines, self._vlines = [], []   # ızgara çizgileri yeni tabloya aittir
        top = ttk.Frame(r, padding=(14, 10, 14, 4))
        top.pack(fill="x")
        ttk.Label(top, text="Next Book", font=("Georgia", 20, "bold"),
                  foreground=t["title_fg"]).pack(side="left")

        # sağ üst: tema noktaları · dil bayrağı · yazı tipi · hakkında
        box = ttk.Frame(top)
        box.pack(side="right")
        self._build_swatches(box)
        self.lang_btn = ttk.Button(box, style="Tool.TButton", text=" ▾", compound="left",
                                   image=self.flag_image(current_language(), 30, 20),
                                   command=self.toggle_lang_popup)
        self.lang_btn.pack(side="left", padx=(12, 0))
        ttk.Button(box, text="Aa", width=3, style="Tool.TButton",
                   command=self.open_font_dialog).pack(side="left", padx=(6, 0))
        ttk.Button(box, text="ⓘ", width=3, style="Tool.TButton",
                   command=self.show_about).pack(side="left", padx=(6, 0))

        # duyuru kutusu: kapatılamaz; üstüne gelince vurgulanır, tıklayınca duyuru panosu açılır.
        # settings.json içinde "announcements": false ise kutu hiç gösterilmez.
        self._ann_over = False
        self.ann_frame = tk.Frame(top, bg=t["bg"], bd=0, highlightthickness=1,
                                  highlightbackground=t["border"], cursor="hand2")
        if self.settings.get("announcements", True):
            self.ann_frame.pack(side="left", fill="x", expand=True, padx=(14, 14), pady=(8, 0))
        self.ann_more = tk.Label(self.ann_frame, text="›", font=self.font_head, cursor="hand2")
        self.ann_more.pack(side="right", padx=(0, 10))
        self.ann_icon = tk.Label(self.ann_frame, bd=0, cursor="hand2")
        self.ann_icon.pack(side="left", padx=(10, 0))
        self.ann_label = tk.Label(self.ann_frame, anchor="w", justify="left", width=1,
                                  font=self.font_normal, cursor="hand2")
        self.ann_label.pack(side="left", fill="x", expand=True, padx=8, pady=3)
        self.ann_label.bind("<Configure>", lambda _e: self._fit_announcement())
        for w in (self.ann_frame, self.ann_icon, self.ann_label, self.ann_more):
            w.bind("<Button-1>", lambda _e: self.open_announcement_board())
            w.bind("<Enter>", self._ann_enter)
            w.bind("<Leave>", self._ann_leave)
        self._render_announcement()

        # arama
        sf = ttk.Frame(r, padding=(14, 4))
        sf.pack(fill="x")
        ttk.Label(sf, text=T("lbl_title")).grid(row=0, column=0, sticky="w")
        ttk.Label(sf, text=T("lbl_author")).grid(row=0, column=1, sticky="w", padx=(8, 0))
        e1 = ttk.Entry(sf, textvariable=self.title_var, width=46)
        e2 = ttk.Entry(sf, textvariable=self.author_var, width=30)
        e1.grid(row=1, column=0, sticky="ew")
        e2.grid(row=1, column=1, sticky="ew", padx=(8, 0))
        attach_autocomplete(e2, lambda: suggest_values(self.all_books(), "author"))
        self.search_btn = ttk.Button(sf, text=T("btn_search"), style="Accent.TButton", command=self.do_search)
        self.search_btn.grid(row=1, column=2, padx=(8, 0))
        ttk.Button(sf, text=T("btn_clear"), command=self.clear_search).grid(row=1, column=3, padx=(6, 0))
        ttk.Button(sf, text=T("btn_manual"), command=self.add_manual).grid(row=1, column=4, padx=(14, 0))
        sf.columnconfigure(0, weight=3)
        sf.columnconfigure(1, weight=2)
        for e in (e1, e2):
            e.bind("<Return>", lambda _e: self.do_search())
        e1.focus_set()

        # arama sonuçları: sadece sonuç varken görünür
        self.res_frame = ttk.LabelFrame(r, text=T("results_title"), padding=6)
        cols = ("title", "author", "genre", "pages", "year")
        self.res_tree = ttk.Treeview(self.res_frame, columns=cols, show="headings", height=6, selectmode="browse")
        for cid, w in (("title", 380), ("author", 230), ("genre", 170), ("pages", 70), ("year", 80)):
            self.res_tree.heading(cid, text=T("col_" + cid))
            self.res_tree.column(cid, width=w, stretch=(cid == "title"),
                                 anchor="center" if cid in ("pages", "year") else "w")
        rs = ttk.Scrollbar(self.res_frame, orient="vertical", command=self.res_tree.yview)
        self.res_tree.configure(yscrollcommand=rs.set)
        self.res_tree.pack(side="left", fill="x", expand=True)
        rs.pack(side="left", fill="y")
        ttk.Button(self.res_frame, text=T("btn_add_to_list"), command=self.add_selected_result
                   ).pack(side="left", padx=(8, 0), anchor="n")
        self.res_tree.tag_configure("added", foreground=t["unscored_fg"])
        self.res_tree.bind("<Double-1>", lambda _e: self.add_selected_result())
        self.res_tree.bind("<Button-1>", lambda e: self._clear_if_blank(self.res_tree, e), add="+")

        # ana liste
        self.tabbar = tk.Frame(r, bg=t["bg"], bd=0, highlightthickness=0)
        self.tabbar.pack(fill="x", padx=14, pady=(6, 0))
        self.tabbar.bind("<Configure>", lambda _e: self._layout_tabs())
        self.main_frame = mf = tk.Frame(r, bg=t["bg"], bd=0, highlightthickness=1,
                                        highlightbackground=t["border"], highlightcolor=t["border"])
        mf.pack(fill="both", expand=True, padx=14, pady=(0, 4))
        self._build_tabs()
        self.tree = ttk.Treeview(mf, show=("tree", "headings"), selectmode="extended")
        vs = ttk.Scrollbar(mf, orient="vertical", command=self.tree.yview)
        hs = ttk.Scrollbar(mf, orient="horizontal", command=self.tree.xview)
        self.vs, self.hs = vs, hs
        self.tree.configure(yscrollcommand=self._yset, xscrollcommand=self._xset)
        self.tree.grid(row=0, column=0, sticky="nsew", padx=(6, 0), pady=(6, 0))
        vs.grid(row=0, column=1, sticky="ns", padx=(0, 6), pady=(6, 0))
        hs.grid(row=1, column=0, sticky="ew", padx=(6, 0), pady=(0, 6))
        mf.rowconfigure(0, weight=1)
        mf.columnconfigure(0, weight=1)
        self.tree.tag_configure("unscored", foreground=t["text"], background=t["tree_bg"], font=self.font_normal)
        for i, col in enumerate(t["ramp"]):
            self.tree.tag_configure(f"s{i}", background=col, foreground=t["scored_fg"], font=self.font_normal)
        self.tree.bind("<Double-1>", self.on_double_click)
        self.tree.bind("<Delete>", lambda _e: self.delete_selected())
        self.tree.bind("<MouseWheel>", lambda _e: self.finish_edit(True))
        self.tree.bind("<Shift-MouseWheel>", self._on_shift_wheel)
        self.tree.bind("<Button-1>", lambda e: self._clear_if_blank(self.tree, e), add="+")
        self.tree.bind("<Motion>", self._on_tree_motion)
        self.tree.bind(RCLICK, self._on_row_menu)
        if sys.platform == "darwin":
            self.tree.bind("<Control-Button-1>", self._on_row_menu)   # Ctrl+tık = sağ tık
        self.tree.bind("<Button-1>", self._on_tree_click, add="+")
        for seq in ("<Configure>", "<B1-Motion>", "<ButtonRelease-1>"):
            self.tree.bind(seq, self._schedule_grid, add="+")
        self._build_status_icons()

        # alt çubuk
        bf = ttk.Frame(r, padding=(14, 4, 14, 4))
        bf.pack(fill="x")
        ttk.Button(bf, text=T("btn_settings"), command=self.open_settings).pack(side="left")
        ttk.Button(bf, text=T("btn_import"), command=self.import_from_excel).pack(side="left", padx=(18, 0))
        ttk.Button(bf, text=T("btn_export"), command=self.export_to_excel).pack(side="left", padx=6)
        self.grid_var = tk.BooleanVar(value=self.grid_on)
        ttk.Checkbutton(bf, text="▦ " + T("btn_grid"), style="Grid.Toolbutton", variable=self.grid_var,
                        command=self.toggle_grid).pack(side="left")
        self.status = ttk.Label(bf, text="", foreground=t["muted"])
        self.status.pack(side="right")
        hint = ttk.Label(r, foreground=t["hint"], padding=(14, 0, 14, 8), text=T("hint") + " · " + T("hint_blank"), justify="left")
        hint.pack(fill="x")
        hint.bind("<Configure>", lambda e: hint.configure(wraplength=max(200, e.width - 28)))

    # ---------------------------------------------------------------- sekmeler (listeler)
    def _tab_h(self):
        return self.font_head.metrics("linespace") + 14

    def _build_tabs(self):
        """Her liste için bir sekme + sonda “+” düğmesi. Ekleme/silme/sıralama/ad değişince yeniden kurulur."""
        for w in self.tabbar.winfo_children():
            w.destroy()
        self.tab_items = []
        for lst in self.store["lists"]:
            fr = tk.Frame(self.tabbar, bd=0, highlightthickness=1, cursor="hand2")
            fr.pack_propagate(False)
            nm = tk.Label(fr, bd=0, anchor="w", font=self.font_head, cursor="hand2")
            nm.pack(side="left", fill="both", expand=True, padx=(10, 0))
            xb = None
            if len(self.store["lists"]) > 1:
                xb = tk.Label(fr, text="×", bd=0, font=self.font_head, cursor="hand2", padx=7)
                xb.pack(side="right", fill="y")
                xb.bind("<Button-1>", lambda _e, l=lst: self.delete_list(l))
                xb.bind("<Enter>", lambda _e, w=xb: w.configure(fg=self.theme["accent"]))
                xb.bind("<Leave>", lambda _e, l=lst: self._style_tabs())
            for w in (fr, nm):
                w.bind("<ButtonPress-1>", lambda e, l=lst: self._tab_press(l, e))
                w.bind("<B1-Motion>", self._tab_motion)
                w.bind("<ButtonRelease-1>", self._tab_release)
                w.bind("<Double-Button-1>", lambda _e, l=lst: self.rename_list(l))
                w.bind(RCLICK, lambda e, l=lst: self._tab_menu(l, e))
                w.bind(MCLICK, lambda _e, l=lst: self.delete_list(l))
                w.bind("<Enter>", lambda _e, l=lst: self._style_tabs(hover=l))
                w.bind("<Leave>", lambda _e: self._style_tabs())
            self.tab_items.append({"lst": lst, "fr": fr, "nm": nm, "x": xb, "pos": (0, 0)})
        plus = tk.Label(self.tabbar, text="+", bd=0, font=self.font_head, cursor="hand2")
        plus.bind("<Button-1>", lambda _e: self.add_list())
        plus.bind("<Enter>", lambda _e: plus.configure(fg=self.theme["accent"]))
        plus.bind("<Leave>", lambda _e: plus.configure(fg=self.theme["text"]))
        self.tab_plus = plus
        self._layout_tabs()
        self._style_tabs()

    def _layout_tabs(self):
        items = getattr(self, "tab_items", None)
        if not items:
            return
        avail = self.tabbar.winfo_width()
        h, gap = self._tab_h(), 3
        self.tabbar.configure(height=h + 1)
        if avail <= 1:
            return
        plus_w = h + 2
        nat = [min(self.font_head.measure(it["lst"]["name"]) + (50 if it["x"] else 30), 260) for it in items]
        room = max(70, (avail - plus_w - gap * (len(items) + 1)) // len(items))
        x = 0
        for it, nw in zip(items, nat):
            w = min(nw, room)
            it["fr"].place(x=x, y=0, width=w, height=h)
            it["pos"] = (x, w)
            it["nm"].configure(text=fit_text(self.font_head, it["lst"]["name"], w - (36 if it["x"] else 18)))
            x += w + gap
        self.tab_plus.place(x=x, y=0, width=plus_w, height=h)

    def _style_tabs(self, hover=None):
        t = self.theme
        for it in getattr(self, "tab_items", []):
            if it["lst"] is self.data:
                bg, fg = t["header_bg"], t["header_fg"]
            elif it["lst"] is hover:
                bg, fg = t["button_active"], t["text"]
            else:
                bg, fg = t["button"], t["text"]
            it["fr"].configure(bg=bg, highlightbackground=t["border"], highlightcolor=t["border"])
            it["nm"].configure(bg=bg, fg=fg)
            if it["x"] is not None:
                it["x"].configure(bg=bg, fg=fg)
        if getattr(self, "tab_plus", None) is not None:
            self.tab_plus.configure(bg=t["bg"], fg=t["text"])

    def _tab_press(self, lst, e):
        self._drag = {"lst": lst, "x0": e.x_root, "moved": False}

    def _tab_motion(self, e):
        d = self._drag
        if not d:
            return
        if abs(e.x_root - d["x0"]) > 6:
            d["moved"] = True
        if d["moved"]:   # sekme fareyi izlesin
            it = next(i for i in self.tab_items if i["lst"] is d["lst"])
            x0, w = it["pos"]
            nx = max(0, min(self.tabbar.winfo_width() - w, x0 + e.x_root - d["x0"]))
            it["fr"].place(x=nx)
            it["fr"].lift()

    def _tab_release(self, e):
        d, self._drag = self._drag, None
        if not d:
            return
        if not d["moved"]:
            self.switch_list(d["lst"])
            return
        lists = self.store["lists"]
        cur = lists.index(d["lst"])
        it = self.tab_items[cur]
        center = it["pos"][0] + it["pos"][1] / 2 + (e.x_root - d["x0"])
        target = sum(1 for j, o in enumerate(self.tab_items)
                     if j != cur and o["pos"][0] + o["pos"][1] / 2 < center)
        if target != cur:
            lists.insert(target, lists.pop(cur))
            self.save()
        self.root.after_idle(self._build_tabs)

    def _tab_menu(self, lst, e):
        t = self.theme
        m = tk.Menu(self.root, tearoff=0, bg=t["field"], fg=t["field_fg"], font=self.font_normal,
                    activebackground=t["select_bg"], activeforeground=t["select_fg"])
        m.add_command(label=T("tab_new"), command=self.add_list)
        m.add_command(label=T("tab_rename"), command=lambda: self.rename_list(lst))
        m.add_command(label=T("tab_duplicate"), command=lambda: self.duplicate_list(lst))
        m.add_separator()
        m.add_command(label=T("tab_close"), command=lambda: self.delete_list(lst),
                      state="normal" if len(self.store["lists"]) > 1 else "disabled")
        try:
            m.tk_popup(e.x_root, e.y_root)
        finally:
            m.grab_release()

    def unique_list_name(self, base, ignore=None):
        base = (base or "").strip()[:56] or T("tab_new")
        taken = {l["name"].casefold() for l in self.store["lists"] if l is not ignore}
        name, n = base, 2
        while name.casefold() in taken:
            name = f"{base} ({n})"
            n += 1
        return name

    def _activate(self, lst):
        """Açık sekmeyi değiştir: tablo, sütunlar ve arama sonuçlarındaki ✓ işaretleri bu listeye göre yenilenir."""
        self.finish_edit(True)
        self.data = lst
        self.store["active"] = lst["id"]
        self.sort_col, self.sort_rev = None, False
        self._hover_cid = None
        self.save()
        self._style_tabs()
        self.rebuild_columns()
        self.refresh_table()
        if self.results:
            self.show_results(self.results)

    def switch_list(self, lst):
        if lst is not self.data:
            self._activate(lst)

    def _limit_reached(self):
        if len(self.store["lists"]) >= MAX_LISTS:
            self.root.bell()
            self.set_status(T("tab_limit", n=MAX_LISTS))
            return True
        return False

    def add_list(self):
        if self._limit_reached():
            return
        self.finish_edit(True)
        name = simpledialog.askstring(APP_NAME, T("ask_list_name"), parent=self.root,
                                      initialvalue=self.unique_list_name(T("tab_new")))
        if not name or not name.strip():
            return
        lst = new_list(self.unique_list_name(name), default_data(current_language()))
        self.store["lists"].append(lst)
        self._build_tabs()
        self._activate(lst)

    def rename_list(self, lst):
        self.finish_edit(True)
        name = simpledialog.askstring(APP_NAME, T("ask_list_name"), parent=self.root, initialvalue=lst["name"])
        if not name or not name.strip() or name.strip() == lst["name"]:
            return
        lst["name"] = self.unique_list_name(name, ignore=lst)
        lst.pop("name_builtin", None)
        self.save()
        self._build_tabs()

    def duplicate_list(self, lst):
        if self._limit_reached():
            return
        self.finish_edit(True)
        import copy
        c = copy.deepcopy(lst)
        c["id"] = new_id()
        c["name"] = self.unique_list_name(lst["name"])
        c.pop("name_builtin", None)
        self.store["lists"].insert(self.store["lists"].index(lst) + 1, c)
        self._build_tabs()
        self._activate(c)

    def delete_list(self, lst):
        lists = self.store["lists"]
        if len(lists) <= 1:
            return
        self.finish_edit(False)
        if lst["books"] and not messagebox.askyesno(
                APP_NAME, T("confirm_delete_list", name=lst["name"], n=len(lst["books"]))):
            return
        idx = lists.index(lst)
        lists.remove(lst)
        self._build_tabs()
        if lst is self.data:
            self._activate(lists[max(0, idx - 1)])
        else:
            self.save()

    # ---------------------------------------------------------------- tema noktaları
    def _swatch_xs(self):
        """Her tema topunun sol x değeri; aydınlık ve karanlık gruplar arasında küçük bir boşluk var."""
        sw, xs, x = self.SWATCH, [], 8
        light = sum(1 for th in THEMES.values() if not th["dark"])
        for i in range(len(THEMES)):
            if i == light:
                x += 14
            xs.append(x)
            x += sw + 8
        return xs

    def _build_swatches(self, parent):
        """Her tema için tek renkli bir top (üstte hafif parlama, seçili olanın etrafında halka)."""
        t, sw = self.theme, self.SWATCH
        xs = self._swatch_xs()
        c = tk.Canvas(parent, width=xs[-1] + sw + 12, height=sw + 12, bg=t["bg"],
                      highlightthickness=0, bd=0, cursor="hand2")
        light = sum(1 for th in THEMES.values() if not th["dark"])
        if 0 < light < len(xs):   # gruplar arası ince ayırıcı
            gx = xs[light] - 11
            c.create_line(gx, 8, gx, sw + 4, fill=t["border"])
        for x0, (name, th) in zip(xs, THEMES.items()):
            y0 = 6
            if name == self.theme_name:
                c.create_oval(x0 - 4, y0 - 4, x0 + sw + 4, y0 + sw + 4, outline=t["accent"], width=2)
            base = th["swatch"]
            c.create_oval(x0, y0, x0 + sw, y0 + sw, fill=base, outline=_mix(base, t["text"], 0.35))
            hl = _mix(base, "#FFFFFF", 0.45)
            c.create_oval(x0 + sw * 0.22, y0 + sw * 0.16, x0 + sw * 0.52, y0 + sw * 0.40, fill=hl, outline="")
        c.bind("<Button-1>", self._on_swatch_click)
        c.pack(side="left")
        self.swatch_canvas = c

    def _on_swatch_click(self, e):
        for x0, name in zip(self._swatch_xs(), THEMES):
            if x0 - 4 <= e.x <= x0 + self.SWATCH + 4:
                self.set_theme(name)
                return

    # ---------------------------------------------------------------- duyurular
    def ann_icon_img(self, kind, fill, glyph, bg, size=None):
        """Duyuru simgesi (önbellekli; PhotoImage'lar bellekte tutulmazsa ekrandan kaybolur)."""
        size = size or max(16, round(self.font_head.metrics("linespace") * 0.95))
        key = (kind, size, fill, glyph, bg)
        if key not in self.icon_cache:
            self.icon_cache[key] = make_ann_icon(kind, size, fill, glyph, bg)
        return self.icon_cache[key]

    def _announce_tick(self):
        """Duyuru dosyasını arka planda çek; açık kaldığı sürece periyodik tekrarla."""
        if ANNOUNCE_URL and self.settings.get("announcements", True):
            threading.Thread(target=self._announce_worker, daemon=True).start()
        self.root.after(ANNOUNCE_REFRESH_MS, self._announce_tick)

    def _announce_worker(self):
        try:
            self.q.put(("ann", fetch_feed(ANNOUNCE_URL)))
        except Exception:
            pass   # çevrimdışı / sunucu yok: sessizce slogan görünmeye devam eder

    def _active_announcements(self):
        if not self.ann_feed or not self.settings.get("announcements", True):
            return []
        return select_announcements(self.ann_feed, APP_ID, APP_VERSION, current_language())

    def _cancel_ann_job(self):
        if self._ann_job is not None:
            try:
                self.root.after_cancel(self._ann_job)
            except tk.TclError:
                pass
            self._ann_job = None

    def _render_announcement(self):
        try:
            if not self.ann_label.winfo_exists():
                return
        except (AttributeError, tk.TclError):
            return
        self._cancel_ann_job()
        items = self._active_announcements()
        if not items:
            self.ann_idx = 0
            self.ann_cur = None
            self.ann_full = T("ann_empty")
            more = "›"
        else:
            self.ann_idx %= len(items)
            cur = self.ann_cur = items[self.ann_idx]
            self.ann_full = cur["text"]
            more = f"{self.ann_idx + 1}/{len(items)}  ›" if len(items) > 1 else "›"
            if len(items) > 1:
                self._ann_job = self.root.after(ANNOUNCE_ROTATE_MS, self._next_announcement)
        try:
            self.ann_more.configure(text=more)
        except tk.TclError:
            pass
        self._paint_announcement()
        self._fit_announcement()

    def _paint_announcement(self):
        """Kutunun renkleri: duyuru varsa vurgulu, yoksa sade; fare üstündeyken biraz daha koyu."""
        t = self.theme
        has = self.ann_cur is not None
        base = t["ramp"][1] if has else t["bg"]
        bg = _mix(base, t["accent"], 0.22) if self._ann_over else base
        edge = t["accent"] if (has or self._ann_over) else t["border"]
        kind = self.ann_cur["type"] if has else "info"
        fill, glyph = ann_icon_colors(t, kind)
        try:
            self.ann_frame.configure(bg=bg, highlightbackground=edge)
            self.ann_icon.configure(bg=bg, image=self.ann_icon_img(kind, fill, glyph, bg))
            self.ann_label.configure(bg=bg, fg=t["text"] if has else t["muted"],
                                     font=self.font_head if has else self.font_normal)
            self.ann_more.configure(bg=bg, fg=t["muted"])
        except (AttributeError, tk.TclError):
            pass

    def _fit_announcement(self):
        try:
            font = tkfont.Font(font=self.ann_label.cget("font"))
            self.ann_label.configure(text=fit_text(font, self.ann_full, self.ann_label.winfo_width() - 16))
        except (AttributeError, tk.TclError):
            pass

    def _next_announcement(self):
        self._ann_job = None
        self.ann_idx += 1
        self._render_announcement()

    # ---- fare üstüne gelince vurgula (alt widget'lar arasında geçişte titremesin diye gecikmeli kontrol)
    def _ann_enter(self, _e=None):
        self._ann_over = True
        self._paint_announcement()

    def _ann_leave(self, _e=None):
        self.root.after(40, self._ann_check_leave)

    def _ann_check_leave(self):
        try:
            w = self.root.winfo_containing(*self.root.winfo_pointerxy())
            frame = str(self.ann_frame)
            inside = w is not None and (str(w) == frame or str(w).startswith(frame + "."))
        except (tk.TclError, KeyError, AttributeError):
            inside = False
        if inside != self._ann_over:
            self._ann_over = inside
            self._paint_announcement()

    def open_announcement_board(self):
        self.finish_edit(True)
        AnnouncementBoard(self, self._active_announcements())

    def open_url_confirmed(self, text, url, parent=None):
        """Duyuru bağlantısı uzaktan geldiği için açmadan önce kullanıcıya sor."""
        if messagebox.askyesno(APP_NAME, text + "\n\n" + T("ann_open_q", url=url), parent=parent or self.root):
            webbrowser.open(url)

    # ---------------------------------------------------------------- dil açılır penceresi
    def toggle_lang_popup(self):
        if self.lang_popup is not None:
            try:
                self.lang_popup.destroy()
            except tk.TclError:
                pass
            self.lang_popup = None
            return
        self.open_lang_popup()

    def open_lang_popup(self):
        t = self.theme
        pop = tk.Toplevel(self.root)
        pop.overrideredirect(True)
        pop.configure(bg=t["border"])
        inner = tk.Frame(pop, bg=t["bg"], padx=8, pady=8)
        inner.pack(padx=1, pady=1)
        fam = self.default_family()
        for i, (code, name) in enumerate(LANGS):
            sel = code == current_language()
            btn = tk.Button(inner, image=self.flag_image(code, 44, 30), text=name, compound="top",
                            font=(fam, 9), bd=0, relief="flat", bg=t["bg"], fg=t["text"],
                            activebackground=t["button_active"], activeforeground=t["text"],
                            highlightthickness=2, highlightbackground=t["accent"] if sel else t["bg"],
                            highlightcolor=t["accent"], padx=8, pady=4, cursor="hand2",
                            command=lambda c=code: self._choose_lang(c))
            btn.grid(row=i // 2, column=i % 2, padx=3, pady=3)
        try:
            self.root.update_idletasks()
            x = self.lang_btn.winfo_rootx()
            y = self.lang_btn.winfo_rooty() + self.lang_btn.winfo_height() + 4
            pop.update_idletasks()
            x = max(0, min(x, self.root.winfo_screenwidth() - pop.winfo_reqwidth() - 8))
            pop.geometry(f"+{x}+{y}")
        except tk.TclError:
            pass
        pop.bind("<Escape>", lambda _e: self.toggle_lang_popup())
        pop.bind("<FocusOut>", lambda _e: self._popup_focus_out(pop))
        self.lang_popup = pop
        pop.focus_force()

    def _popup_focus_out(self, pop):
        def check():
            if self.lang_popup is not pop:
                return
            try:
                f = self.root.focus_get()
            except Exception:
                f = None
            if f is None or not str(f).startswith(str(pop)):
                try:
                    pop.destroy()
                except tk.TclError:
                    pass
                self.lang_popup = None
        self.root.after(80, check)

    def _choose_lang(self, code):
        if self.lang_popup is not None:
            try:
                self.lang_popup.destroy()
            except tk.TclError:
                pass
            self.lang_popup = None
        self.set_lang(code)

    def open_font_dialog(self):
        self.finish_edit(True)
        FontDialog(self)

    def show_about(self):
        AboutDialog(self)

    # ---------------------------------------------------------------- seçimi temizleme
    def _clear_if_blank(self, tree, event):
        """Tablonun boş bir yerine tıklanınca hiçbir satır seçili kalmasın."""
        region = tree.identify_region(event.x, event.y)
        if region in ("heading", "separator"):
            return
        if not tree.identify_row(event.y):
            tree.selection_remove(tree.selection())

    def _on_root_click(self, event):
        try:
            cls = event.widget.winfo_class()
        except Exception:
            return
        if cls in ("TFrame", "TLabel", "TLabelframe", "Frame", "Label", "Tk"):
            for tr in (getattr(self, "tree", None), getattr(self, "res_tree", None)):
                if tr is not None:
                    try:
                        tr.selection_remove(tr.selection())
                    except tk.TclError:
                        pass

    # ---------------------------------------------------------------- sütunlar
    def col_ids(self):
        return self.FIXED + ["c_" + c["id"] for c in self.data["criteria"]] + ["score"]

    def criterion_by_col(self, cid):
        return next((c for c in self.data["criteria"] if "c_" + c["id"] == cid), None)

    def heading_text(self, cid):
        if cid in self.FIXED or cid == "score":
            return T("col_" + cid)
        c = self.criterion_by_col(cid)
        return short(c["name"], 22) if c else cid

    def rebuild_columns(self):
        ids = self.col_ids()
        # Önce görünür sütun listesini sıfırla: silinen ölçütün eski sütununa
        # başvuru kalırsa Tk "Invalid column index" hatası verir.
        self.tree["displaycolumns"] = "#all"
        self.tree["columns"] = ids
        self.tree["displaycolumns"] = ids
        self.update_headings()
        for cid in ids:
            if cid == "title":
                self.tree.column(cid, width=320, minwidth=200, stretch=True, anchor="w")
            elif cid == "author":
                self.tree.column(cid, width=200, minwidth=100, stretch=False, anchor="w")
            elif cid == "genre":
                self.tree.column(cid, width=130, minwidth=80, stretch=False, anchor="w")
            elif cid == "pages":
                self.tree.column(cid, width=70, minwidth=50, stretch=False, anchor="center")
            elif cid == "score":
                w = self.font_head.measure(self.heading_text(cid)) + 46
                self.tree.column(cid, width=max(110, w), minwidth=90, stretch=False, anchor="center")
            else:   # ölçüt sütunu: başlığa göre dar; ne kadar çok olursa yatay kaydırma devreye girer
                w = self.font_head.measure(self.heading_text(cid)) + 46
                self.tree.column(cid, width=max(84, min(200, w)), minwidth=60, stretch=False, anchor="center")
        self.tree.xview_moveto(0)
        self._schedule_grid()

    def autosize_columns(self):
        """Kitap adı / yazar / tür sütunlarını içeriğe göre ayarla (kesilmesin)."""
        books = self.data["books"]
        f = self.font_normal

        def best(key, lo, hi):
            w = max((f.measure(str(b.get(key) or "")) for b in books), default=0)
            return max(lo, min(hi, w + 30))

        self.tree.column("title", width=best("title", 260, 560))
        self.tree.column("author", width=best("author", 150, 280))
        self.tree.column("genre", width=best("genre", 110, 190))

    def update_headings(self):
        for cid in self.col_ids():
            text = self.heading_text(cid)
            if cid == self.sort_col:
                text += " ▼" if self.sort_rev else " ▲"
            self.tree.heading(cid, text=text, command=lambda c=cid: self.sort_by(c))

    def _on_tree_motion(self, event):
        """Kısaltılmış ölçüt başlığının üzerine gelince tam adı durum çubuğunda göster."""
        cid = None
        if self.tree.identify_region(event.x, event.y) == "heading":
            try:
                idx = int(self.tree.identify_column(event.x)[1:]) - 1
                ids = self.col_ids()
                cid = ids[idx] if 0 <= idx < len(ids) else None
            except ValueError:
                cid = None
        if cid == self._hover_cid:
            return
        self._hover_cid = cid
        c = self.criterion_by_col(cid) if cid else None
        if c:
            self.set_status(c["name"])
        else:
            self.update_status()

    def _on_shift_wheel(self, event):
        self.finish_edit(True)
        self.tree.xview_scroll(-3 if event.delta > 0 else 3, "units")

    # ---------------------------------------------------------------- tablo
    def row_values(self, b):
        cs = self.data["criteria"]
        vals = [b["title"], b["author"], b["genre"], "" if b["pages"] in (None, "") else b["pages"]]
        vals += [fmt_num(b["scores"].get(c["id"])) for c in cs]
        vals.append(f"{compute_score(b, cs):.2f}" if is_scored(b, cs) else "—")
        return vals

    def row_tag(self, b):
        cs = self.data["criteria"]
        if not is_scored(b, cs):
            return ("unscored",)
        return (f"s{score_bucket(compute_score(b, cs), self.data['scale_max'], len(self.theme['ramp']))}",)

    def refresh_table(self):
        self.finish_edit(False)
        self.tree.delete(*self.tree.get_children())
        for b in self.data["books"]:
            self.tree.insert("", "end", iid=b["id"], values=self.row_values(b), tags=self.row_tag(b),
                             image=self._row_image(b))
        # en altta her zaman boş bir satır: herhangi bir hücresine yazılınca yeni kitap eklenir
        self.tree.insert("", "end", iid=NEW_IID, values=[""] * len(self.col_ids()), tags=("unscored",))
        self.autosize_columns()
        self.update_status()
        self._schedule_grid()

    def update_row(self, b):
        self.tree.item(b["id"], values=self.row_values(b), tags=self.row_tag(b), image=self._row_image(b))
        self.update_status()

    def update_status(self):
        cs = self.data["criteria"]
        total = len(self.data["books"])
        scored = sum(1 for b in self.data["books"] if is_scored(b, cs))
        self.set_status(T("status_counts", total=total, scored=scored, pending=total - scored))

    def set_status(self, text):
        self.status.config(text=text)

    def book_by_id(self, bid):
        return next((b for b in self.data["books"] if b["id"] == bid), None)

    def save(self):
        try:
            save_data(self.store)
        except OSError as e:
            messagebox.showerror(APP_NAME, T("save_failed", err=e))

    # ---------------------------------------------------------------- sıralama
    def sort_by(self, cid):
        if self.sort_col == cid:
            self.sort_rev = not self.sort_rev
        else:
            self.sort_col = cid
            self.sort_rev = cid == "score" or cid.startswith("c_")   # puanlar: yüksekten düşüğe
        self.apply_sort()

    def apply_sort(self):
        self.finish_edit(True)
        cs = self.data["criteria"]
        col = self.sort_col

        def value(b):
            if col == "score":
                return compute_score(b, cs) if is_scored(b, cs) else None
            if col == "pages":
                return b["pages"] if b["pages"] not in (None, "") else None
            if col.startswith("c_"):
                v = b["scores"].get(col[2:])
                return None if v in (None, "") else float(v)
            t = (b.get(col) or "").strip()
            return t.casefold() if t else None

        present = [b for b in self.data["books"] if value(b) is not None]
        missing = [b for b in self.data["books"] if value(b) is None]   # boşlar hep sonda
        present.sort(key=value, reverse=self.sort_rev)
        self.data["books"] = present + missing
        self.save()
        self.update_headings()
        self.refresh_table()

    # ---------------------------------------------------------------- arama
    def do_search(self):
        t, a = self.title_var.get().strip(), self.author_var.get().strip()
        if not t and not a:
            self.set_status(T("enter_query"))
            return
        self.set_status(T("searching"))
        self.search_btn.config(state="disabled")
        threading.Thread(target=self._search_worker, args=(t, a), daemon=True).start()

    def _search_worker(self, t, a):
        try:
            self.q.put(("ok", search_books(t, a)))
        except Exception as e:   # ağ hataları dahil
            self.q.put(("err", e))

    def _poll_queue(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == "ann":
                    self.ann_feed = payload
                    self._render_announcement()
                    continue
                try:
                    self.search_btn.config(state="normal")
                except tk.TclError:
                    pass
                if kind == "ok":
                    self.show_results(payload)
                else:
                    if isinstance(payload, (urllib.error.URLError, TimeoutError, OSError)):
                        self.set_status(T("err_network"))
                    else:
                        self.set_status(T("err_search", err=payload))
        except queue.Empty:
            pass
        self.root.after(100, self._poll_queue)

    def show_results(self, results):
        self.results = results
        self.res_tree.delete(*self.res_tree.get_children())
        if not results:
            self.res_frame.pack_forget()
            self.set_status(T("no_results"))
            return
        for i, r in enumerate(results):
            added = self.find_duplicate(r) is not None
            self.res_tree.insert("", "end", iid=str(i), tags=("added",) if added else (),
                                 values=(("✓ " if added else "") + r["title"], r["author"], r["genre"],
                                         r["pages"] or "", r["year"]))
        self.res_frame.pack(fill="x", padx=14, pady=(8, 4), before=self.tabbar)   # sekmelerin üstünde
        self.set_status(T("results_found", n=len(results)))

    def clear_search(self):
        self.title_var.set("")
        self.author_var.set("")
        self.results = []
        self.res_tree.delete(*self.res_tree.get_children())
        self.res_frame.pack_forget()
        self.update_status()

    # ---------------------------------------------------------------- ekleme / silme
    def find_duplicate(self, r):
        tkey = (r["title"].casefold(), (r["author"] or "").casefold())
        for b in self.data["books"]:
            if r.get("key") and b.get("key") == r["key"]:
                return b
            if (b["title"].casefold(), (b["author"] or "").casefold()) == tkey:
                return b
        return None

    def add_selected_result(self):
        sel = self.res_tree.selection()
        if not sel:
            self.set_status(T("pick_result"))
            return
        r = self.results[int(sel[0])]
        dup = self.find_duplicate(r)
        if dup:
            self.tree.selection_set(dup["id"])
            self.tree.see(dup["id"])
            messagebox.showinfo(APP_NAME, T("dup_book", title=dup["title"]))
            return
        self._append_book(new_book(r["title"], r["author"], r["genre"], r["pages"], r["key"]))
        self.res_tree.item(sel[0], tags=("added",))

    def add_manual(self):
        dlg = ManualBookDialog(self)
        self.root.wait_window(dlg)
        r = dlg.result
        if r:
            self._append_book(new_book(r["title"], r["author"], r["genre"], r["pages"]))

    def _append_book(self, b):
        self.data["books"].append(b)
        self.save()
        self.tree.insert("", "end", iid=b["id"], values=self.row_values(b), tags=self.row_tag(b),
                         image=self._row_image(b))
        if self.tree.exists(NEW_IID):
            self.tree.move(NEW_IID, "", "end")   # boş satır hep en altta kalsın
        self.autosize_columns()
        self._schedule_grid()
        self.tree.selection_set(b["id"])
        self.tree.see(b["id"])
        self.update_status()

    def _build_status_icons(self):
        """Okuyorum / okudum rozetleri ve #0 (rozet) sütunu. Tema ya da yazı boyutu değişince yeniden kurulur."""
        t = self.theme
        try:
            rh = int(ttk.Style(self.root).lookup("Treeview", "rowheight") or 24)
        except (tk.TclError, ValueError):
            rh = 24
        size = max(14, rh - 10)
        iw = size + 8   # görsel rozetten geniş: rozet görselin ortasında, sütun da görselin etrafında simetrik
        colw = iw + 2 * self.ICON_PAD
        self.status_icons = {
            "reading": make_status_icon("reading", size, t["accent"], t["accent_fg"], ring=t["accent_fg"], width=iw),
            "paused": make_status_icon("paused", size, t["muted"], t["tree_bg"], ring=t["tree_bg"], width=iw),
            "done": make_status_icon("done", size, t["header_bg"], t["header_fg"], ring=t["header_fg"], width=iw),
        }
        self.tree.heading("#0", text="")
        self.tree.column("#0", width=colw, minwidth=colw, stretch=False, anchor="center")

    def _row_image(self, b):
        return getattr(self, "status_icons", {}).get(book_status(b), "")

    def _on_row_menu(self, event):
        """Satıra sağ tık: durum (okunacak / okuyorum / okudum) ve listeden silme. Çoklu seçimde hepsine uygulanır."""
        iid = self.tree.identify_row(event.y)
        if not iid or iid == NEW_IID or self.tree.identify_region(event.x, event.y) == "heading":
            return
        self.finish_edit(True)
        if iid not in self.tree.selection():
            self.tree.selection_set(iid)
        ids = [i for i in self.tree.selection() if i != NEW_IID]
        states = {book_status(b) for b in map(self.book_by_id, ids) if b}
        self._menu_var = tk.StringVar(value=next(iter(states)) if len(states) == 1 else "")
        t = self.theme
        m = tk.Menu(self.root, tearoff=0, bg=t["field"], fg=t["field_fg"], font=self.font_normal,
                    activebackground=t["select_bg"], activeforeground=t["select_fg"],
                    selectcolor=t["field_fg"])
        for code in STATUSES:
            m.add_radiobutton(label=T("st_" + code), variable=self._menu_var, value=code,
                              command=lambda c=code: self.set_books_status(ids, c))
        m.add_separator()
        m.add_command(label=T("row_delete"), command=self.delete_selected)
        try:
            m.tk_popup(event.x_root, event.y_root)
        finally:
            m.grab_release()

    def set_books_status(self, ids, state):
        for bid in ids:
            b = self.book_by_id(bid)
            if not b:
                continue
            if state == "todo":
                b.pop("status", None)
            else:
                b["status"] = state
            self.update_row(b)
        self.save()

    def delete_selected(self):
        sel = tuple(i for i in self.tree.selection() if i != NEW_IID)
        if not sel:
            return
        if not messagebox.askyesno(APP_NAME, T("confirm_delete", n=len(sel))):
            return
        self.finish_edit(False)
        ids = set(sel)
        self.data["books"] = [b for b in self.data["books"] if b["id"] not in ids]
        self.save()
        self.tree.delete(*sel)
        self.update_status()

    # ---------------------------------------------------------------- hücre düzenleme
    def on_double_click(self, event):
        if self.tree.identify_region(event.x, event.y) != "cell":
            return
        iid = self.tree.identify_row(event.y)
        idx = int(self.tree.identify_column(event.x)[1:]) - 1
        if iid:
            self.start_edit(iid, idx)

    def _ensure_col_visible(self, idx):
        """Çok ölçüt varsa düzenlenecek sütun yatay olarak görünür alana kaydırılsın."""
        ids = self.col_ids()
        widths = [int(self.tree.column("#0", "width"))] + [int(self.tree.column(c, "width")) for c in ids]
        total = sum(widths)
        if total <= 0:
            return
        left = sum(widths[:idx + 1])
        right = left + widths[idx + 1]
        first, last = self.tree.xview()
        vis_left, vis_right = first * total, last * total
        if left < vis_left:
            self.tree.xview_moveto(left / total)
        elif right > vis_right:
            self.tree.xview_moveto(max(0.0, (right - (vis_right - vis_left)) / total))
        self.tree.update_idletasks()

    def start_edit(self, iid, idx):
        self.finish_edit(True)
        ids = self.col_ids()
        if idx < 0 or idx >= len(ids) or ids[idx] == "score":
            return
        cid = ids[idx]
        self.tree.see(iid)
        self._ensure_col_visible(idx)
        self.tree.update_idletasks()
        bbox = self.tree.bbox(iid, f"#{idx + 1}")
        if not bbox:
            return
        x, y, w, h = bbox
        b = self.book_by_id(iid)   # boş satırda (NEW_IID) kitap yoktur
        if b is None:
            cur = ""
        elif cid.startswith("c_"):
            cur = fmt_num(b["scores"].get(cid[2:]))
        else:
            cur = "" if b[cid] in (None, "") else str(b[cid])
        t = self.theme
        entry = tk.Entry(self.tree, font=self.font_normal, relief="solid", bd=1,
                         bg=t["field"], fg=t["field_fg"], insertbackground=t["field_fg"],
                         justify="center" if cid.startswith("c_") or cid == "pages" else "left")
        entry.insert(0, cur)
        entry.select_range(0, "end")
        entry.place(x=x, y=y, width=w, height=h)
        entry.focus_set()
        if cid in ("author", "genre"):   # daha önce yazılanları öner
            vals = suggest_values(self.all_books(), cid)
            attach_autocomplete(entry, lambda v=vals: v)
        self.edit = {"iid": iid, "idx": idx, "cid": cid, "entry": entry}
        entry.bind("<Return>", lambda _e: self._commit_and_move(1, 0))
        entry.bind("<Tab>", lambda _e: self._commit_and_move(0, 1))
        entry.bind("<Escape>", lambda _e: self.finish_edit(False))
        entry.bind("<FocusOut>", lambda _e: self.finish_edit(True))

    def _commit_and_move(self, d_row, d_col):
        e = self.edit
        if not e:
            return "break"
        iid, idx = e["iid"], e["idx"]
        nxt = self.tree.next(iid) if self.tree.exists(iid) else ""
        self._created = None
        self.finish_edit(True)
        if iid != NEW_IID and not self.tree.exists(iid):   # satır boşaldığı için silindi
            if nxt and self.tree.exists(nxt):
                self.tree.selection_set(nxt)
                self.root.after(10, lambda: self.start_edit(nxt, idx if d_row else 0))
            return "break"
        if iid == NEW_IID and self._created:   # boş satırda yazılan kitap oluştu: imleç onun satırından devam etsin
            iid = self._created
        ids = self.col_ids()
        nidx = idx + d_col
        niid = iid
        if d_row:
            niid = self.tree.next(iid)
        elif nidx >= len(ids) - 1:   # skor sütunu düzenlenemez → alt satırın ilk ölçütüne geç
            niid = self.tree.next(iid)
            nidx = len(self.FIXED)
        if niid and nidx >= 0:
            self.tree.selection_set(niid)
            self.root.after(10, lambda: self.start_edit(niid, nidx))
        return "break"

    def finish_edit(self, commit=True):
        e = self.edit
        if not e:
            return
        self.edit = None
        text = e["entry"].get().strip()
        try:
            e["entry"].destroy()
        except tk.TclError:
            pass
        if commit:
            self.apply_edit(e["iid"], e["cid"], text)

    def _apply_to_book(self, b, cid, text):
        """Hücre metnini kitaba işler. Geçersizse False döner (kaydetme / satır güncelleme çağıranda)."""
        if cid == "title":
            b["title"] = text   # boş olabilir: yalnızca yazarı olan satırlar geçerli
        elif cid in ("author", "genre"):
            b[cid] = text
        elif cid == "pages":
            if text == "":
                b["pages"] = None
            else:
                try:
                    n = int(float(text.replace(",", ".")))
                    if n < 0:
                        raise ValueError
                    b["pages"] = n
                except (ValueError, OverflowError):
                    self.root.bell()
                    self.set_status(T("bad_pages"))
                    return False
        else:   # puan
            key = cid[2:]
            if text == "":
                b["scores"].pop(key, None)
            else:
                smax = self.data["scale_max"]
                try:
                    v = parse_float(text)
                    if not 1 <= v <= smax:
                        raise ValueError
                except ValueError:
                    self.root.bell()
                    self.set_status(T("bad_score", max=smax))
                    return False
                b["scores"][key] = int(v) if v == int(v) else round(v, 1)
        return True

    @staticmethod
    def is_blank_book(b):
        """Kitap adı, yazar, tür ve sayfa sayısının dördü de boşsa satır boştur (boşluk karakteri de boş sayılır)."""
        return (not (b.get("title") or "").strip() and not (b.get("author") or "").strip()
                and not (b.get("genre") or "").strip() and b.get("pages") in (None, ""))

    def _remove_book_row(self, b):
        self.data["books"] = [x for x in self.data["books"] if x["id"] != b["id"]]
        self.save()
        if self.tree.exists(b["id"]):
            self.tree.delete(b["id"])
        self.update_status()
        self._schedule_grid()

    def apply_edit(self, iid, cid, text):
        self._created = None
        if iid == NEW_IID:   # boş satıra yazıldı → yeni kitap
            if not text:
                return
            b = new_book("")
            if not self._apply_to_book(b, cid, text):
                return
            self._append_book(b)
            self._created = b["id"]
            return
        b = self.book_by_id(iid)
        if not b or not self._apply_to_book(b, cid, text):
            return
        if self.is_blank_book(b):   # dört alan da boşaldı → satırı doğrudan sil
            self._remove_book_row(b)
            return
        self.save()
        self.update_row(b)

    # ---------------------------------------------------------------- ızgara
    # Treeview hücre çizgisi çizemediği için tabloya 1 piksellik ince çubuklar yerleştirilir;
    # yalnızca ekranda görünen satır ve sütunlar için çizilir, kaydırma / boyut değişiminde yenilenir.
    def toggle_grid(self):
        self.grid_on = bool(self.grid_var.get())
        self._persist_settings()
        self._schedule_grid()

    def _yset(self, a, b):
        self.vs.set(a, b)
        self._schedule_grid()

    def _xset(self, a, b):
        self.hs.set(a, b)
        self._schedule_grid()

    def _schedule_grid(self, *_):
        if self._grid_job is None:
            try:
                self._grid_job = self.root.after_idle(self._redraw_grid)
            except tk.TclError:
                self._grid_job = None

    def _forward(self, e, seq):
        """İnce çizginin üstüne denk gelen tıklama / tekerlek olayını tabloya ilet."""
        tr = self.tree
        kw = {"x": e.x_root - tr.winfo_rootx(), "y": e.y_root - tr.winfo_rooty()}
        if seq == "<MouseWheel>":
            kw["delta"] = e.delta
        try:
            tr.event_generate(seq, **kw)
        except tk.TclError:
            pass
        return "break"

    def _grid_line(self, pool, i):
        if i < len(pool):
            return pool[i]
        f = tk.Frame(self.tree, bd=0, highlightthickness=0, bg=self.theme["grid"])
        for src, dst in (("<ButtonPress-1>", "<ButtonPress-1>"), ("<Double-ButtonPress-1>", "<ButtonPress-1>"),
                         ("<ButtonRelease-1>", "<ButtonRelease-1>"), (RCLICK_PRESS, RCLICK_PRESS),
                         ("<MouseWheel>", "<MouseWheel>")):
            f.bind(src, lambda e, d=dst: self._forward(e, d))
        pool.append(f)
        return f

    def _redraw_grid(self):
        self._grid_job = None
        tree = self.tree
        try:
            if not tree.winfo_exists():
                return
        except tk.TclError:
            return
        used_h = used_v = 0
        if self.grid_on:
            try:
                W, H = tree.winfo_width(), tree.winfo_height()
                kids = tree.get_children()
                n = len(kids)
                bb = None
                if n and W > 1 and H > 1:
                    first = min(n - 1, max(0, int(round(tree.yview()[0] * n))))
                    bb = tree.bbox(kids[first])
                if bb:
                    y0, rh = bb[1], bb[3]
                    rows = min(n - first, -(-(H - y0) // rh)) if rh > 0 else 0
                    y_end = min(H, y0 + rows * rh)
                    for k in range(rows):
                        y = y0 + (k + 1) * rh - 1
                        if y >= H:
                            break
                        f = self._grid_line(self._hlines, used_h)
                        used_h += 1
                        f.place(x=0, y=y, width=W, height=1)
                    if rows:
                        cols = ["#0"] + self.col_ids()
                        widths = [int(tree.column(c, "width")) for c in cols]
                        total = sum(widths)
                        x = -round(tree.xview()[0] * total)
                        for w in widths:
                            x += w
                            if 0 < x - 1 < W - 1:
                                f = self._grid_line(self._vlines, used_v)
                                used_v += 1
                                f.place(x=x - 1, y=y0, width=1, height=y_end - y0)
            except tk.TclError:
                pass
        for f in self._hlines[used_h:] + self._vlines[used_v:]:
            f.place_forget()
        if self.edit:   # açık hücre düzenleyici çizgilerin üstünde kalsın
            try:
                self.edit["entry"].lift()
            except tk.TclError:
                pass

    def _on_tree_click(self, event):
        """Boş satırın bir hücresine tek tıklamak doğrudan yazmaya başlatır."""
        if self.tree.identify_region(event.x, event.y) != "cell" or self.tree.identify_row(event.y) != NEW_IID:
            return
        try:
            idx = int(self.tree.identify_column(event.x)[1:]) - 1
        except ValueError:
            return
        if idx >= 0:
            self.root.after(20, lambda: self.tree.exists(NEW_IID) and self.start_edit(NEW_IID, idx))

    # ---------------------------------------------------------------- puanlama ayarları
    def open_settings(self):
        self.finish_edit(True)
        SettingsDialog(self)

    def apply_settings(self, criteria, scale_max):
        alive = {c["id"] for c in criteria}
        for b in self.data["books"]:
            b["scores"] = {k: min(v, scale_max) for k, v in b["scores"].items() if k in alive}
        self.data["criteria"] = criteria
        self.data["scale_max"] = scale_max
        if self.sort_col and self.sort_col.startswith("c_") and self.sort_col[2:] not in alive:
            self.sort_col = None
        self.save()
        self.rebuild_columns()
        self.refresh_table()

    # ---------------------------------------------------------------- Excel
    def import_from_excel(self):
        """Excel dosyası mevcut listenin üstüne yazılmaz; yeni bir sekme olarak eklenir."""
        if not _require_openpyxl():
            return
        if self._limit_reached():
            return
        path = filedialog.askopenfilename(title=T("pick_excel"),
                                          filetypes=[("Excel", "*.xlsx *.xlsm"), ("*", "*.*")])
        if not path:
            return
        try:
            new = import_excel(path)
        except Exception as e:
            messagebox.showerror(APP_NAME, T("import_failed", err=e))
            return
        lst = new_list(self.unique_list_name(Path(path).stem), new)
        self.store["lists"].append(lst)
        self._build_tabs()
        self._activate(lst)
        messagebox.showinfo(APP_NAME, T("import_done_tab", name=lst["name"],
                                        books=len(lst["books"]), crit=len(lst["criteria"])))

    def export_to_excel(self):
        if not _require_openpyxl():
            return
        path = filedialog.asksaveasfilename(title=T("save_excel"), defaultextension=".xlsx",
                                            initialfile=(re.sub(r'[\\/:*?"<>|]+', "_", self.data["name"]).strip() or "Next_Book") + ".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        try:
            export_excel(self.data, path)
        except PermissionError:
            messagebox.showerror(APP_NAME, T("export_locked"))
            return
        except Exception as e:
            messagebox.showerror(APP_NAME, T("export_failed", err=e))
            return
        self.set_status(T("saved_to", path=path))

    # ---------------------------------------------------------------- kapanış
    def on_close(self):
        self.finish_edit(True)
        self.save()
        self._persist_settings()
        self.root.destroy()


def main():
    try:   # Windows'ta net yazı için
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    try:   # görev çubuğunda Python değil Next Book simgesi görünsün
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("NextBook.App")
    except Exception:
        pass
    root = tk.Tk()
    NextBookApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
